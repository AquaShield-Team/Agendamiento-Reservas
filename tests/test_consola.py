# -*- coding: utf-8 -*-
"""Modo consola y utilitarios: main(), ejecutar_login, ejecutar_reservas,
Registro, pausa_manual y los argumentos de Chrome. Con navieras falsas y
Playwright falso, dentro de la sandbox.

Rareza fotografiada: ejecutar_login abre el navegador aunque ninguna naviera
pedida tenga credenciales. (Hasta d86fbca, ejecutar_reservas escribía «SIN EMITIR»
aunque la reserva volviera EMITIDA; ahora solo con OK-EJEMPLO.)
"""
import ast
import contextlib
import io
import os
import re
import shutil
import stat
import sys
import tkinter.messagebox
import types
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

import sinteticos
import soporte


class Llamadas:
    def __init__(self):
        self.lista = []

    def __call__(self, nombre, retorno=None):
        def f(*a, **k):
            self.lista.append((nombre, a, {x: y for x, y in k.items() if x != "esperar_cierre"},
                               "esperar_cierre" in k and k["esperar_cierre"] is not None))
            return retorno
        return f


class TestMain(soporte.CasoAQ):
    def setUp(self):
        super().setUp()
        self.orig = {n: getattr(self.mod, n) for n in ("lanzar_web", "lanzar_panel", "ejecutar_login", "ejecutar_reservas")}
        self.ll = Llamadas()
        for n in self.orig:
            setattr(self.mod, n, self.ll(n))
        os.environ.pop("AQUASHIELD_AUTOCLOSE", None)

    def tearDown(self):
        for n, f in self.orig.items():
            setattr(self.mod, n, f)
        super().tearDown()

    def main(self, *args):
        self.ll.lista.clear()
        sys.argv[:] = ["AQUASHIELD.py", *args]
        try:
            self.mod.main()
        finally:
            soporte.apagar_llaves()
        return self.ll.lista

    def test_despacho_de_main(self):
        todas = list(sinteticos.NAVIERAS)
        self.assertEqual(self.main(), [("lanzar_web", (), {}, False)])
        self.assertEqual(self.main("", "  "), [("lanzar_web", (), {}, False)])
        self.assertEqual(self.main("panel"), [("lanzar_panel", (), {}, False)])
        self.assertEqual(self.main("ventana"), [("lanzar_panel", (), {}, False)])
        self.assertEqual(self.main("Op_Prueba", "ONE"), [("ejecutar_login", ("op_prueba", ["one"]), {}, False)])
        self.assertEqual(self.main("op", "todas"), [("ejecutar_login", ("op", todas), {}, False)])
        self.assertEqual(self.main("op", "xx"), [])
        self.assertEqual(self.main("op", "cma", "reservas"), [("ejecutar_reservas", ("op", "cma"), {}, False)])
        self.assertEqual(self.main("op", "cma", "reservas", "auto"), [("ejecutar_reservas", ("op", "cma"), {}, True)])
        self.assertEqual(self.main("op", "todas", "reserva"),
                         [("ejecutar_reservas", ("op", n), {}, False) for n in todas])

    def test_si_el_panel_web_falla_abre_la_ventana(self):
        avisos = []
        orig = tkinter.messagebox.showwarning
        tkinter.messagebox.showwarning = lambda t, m: avisos.append((t, m))

        def falla(*a, **k):
            raise OSError("puerto ocupado de prueba")
        self.mod.lanzar_web = falla
        try:
            llamadas = self.main()
        finally:
            tkinter.messagebox.showwarning = orig
        self.assertEqual(avisos, [("AQUASHIELD", "No pude abrir el index en el navegador:\npuerto ocupado de prueba\n\n"
                                                 "Abro la ventana clásica.")])
        self.assertEqual(llamadas, [("lanzar_panel", (), {}, False)])


RESPALDOS = "Reservas AQUASHIELD - respaldo *.xlsx"


class ConPlanilla(soporte.CasoAQ):
    def setUp(self):
        super().setUp()
        for p in self.sb.glob(RESPALDOS):          # cada prueba cuenta solo sus respaldos
            p.unlink()
        self.planilla = sinteticos.crear_planilla(self.sb / "Reservas AQUASHIELD.xlsx")
        os.environ.pop("AQUASHIELD_SOLO_PRIMERA", None)
        os.environ.pop("AQUASHIELD_TRAZA", None)


class TestEjecutarReservas(ConPlanilla):
    def test_reservas_por_consola(self):
        def respuesta(nav, rsv, on_pausa):
            return ("EMITIDA", "Booking: FALSO777") if rsv["fila"] == 6 else ("OK-EJEMPLO", "armada falsa")
        antes = {p.name for p in (self.sb / "logs").glob("*")}
        with soporte.Navieras(self.mod, respuesta=respuesta) as n:
            res, carpeta = self.mod.ejecutar_reservas("op_prueba", "cma", esperar_cierre=lambda: None)
        self.assertEqual(res, {5: "OK-EJEMPLO", 6: "EMITIDA"})
        self.assertEqual([kw["chromium_sandbox"] for kw in n.pw.lanzamientos], [True])     # el sandbox (encargo 53)
        self.assertEqual([r[:2] for r in n.reservas], [("cma", 5), ("cma", 6)])
        self.assertEqual(Path(carpeta).parent, self.sb / "logs")
        # Con «_2», «_3»… si otra corrida ya tomó ese segundo; la carpeta tiene que ser nueva (hasta 393f504
        # una corrida del mismo segundo escribía en la carpeta de otra, y esta prueba no lo veía).
        self.assertRegex(Path(carpeta).name, r"^reservas_cma_op_prueba_\d{8}_\d{6}(_\d+)?$")
        self.assertNotIn(Path(carpeta).name, antes)
        ws = openpyxl.load_workbook(self.planilla)["CMA-CGM"]
        patron = (r"^SOLICITUD {i} · {e} · nave: {nave} · PUERTO ALFA→{dest} · \d\d-\d\d \d\d:\d\d · "
                  r"ref\(ej\) EJ\d{{8}} · {det}{cola}$")
        self.assertRegex(ws["H5"].value, patron.format(i=1, e="OK-EJEMPLO", nave="NAVE PRUEBA CINCO",
                                                       dest="CIUDAD GAMMA", det="armada falsa", cola=" · SIN EMITIR"))
        # EMITIDA: el texto dice lo que volvió, sin «SIN EMITIR» (hasta d86fbca lo agregaba siempre).
        self.assertRegex(ws["H6"].value, patron.format(i=2, e="EMITIDA", nave="NAVE PRUEBA SEIS",
                                                       dest="CIUDAD KAPPA", det="Booking: FALSO777", cola=""))

    def test_sin_emitir_solo_si_volvio_sin_emitir(self):
        # Solo OK-EJEMPLO lleva «SIN EMITIR». (Hasta f03bb75 un ERROR podía venir después del
        # botón final; ahora lo de después de la guarda es EMITIDA, ENVIADA – REVISAR EN PORTAL o
        # NO ENVIADA: test_estados_tras_el_envio_en_la_planilla.)
        estados = {5: ("OK-EJEMPLO", "armada falsa"), 6: ("EMITIDA", "Booking: FALSO778"),
                   7: ("REVISAR", "revisar falso"), 10: ("ERROR", "error falso")}
        self.poner_nave("ONE", 10, "NAVE PRUEBA TRECE")      # «NAVE POR COMPLETAR» ya no va al portal (CICLO-solo-la-nave-pedida.md)
        self.poner_nave("ONE", 7, "NAVE PRUEBA CATORCE")     # la columna F ya no es la nave (CICLO-cola-seis-items.md)
        with soporte.Navieras(self.mod, respuesta=lambda nav, rsv, on_pausa: estados[rsv["fila"]]):
            res, _ = self.mod.ejecutar_reservas("op_prueba", "one", esperar_cierre=lambda: None)
        self.assertEqual(res, {fila: e for fila, (e, _) in estados.items()})
        ws = openpyxl.load_workbook(self.planilla)["ONE"]
        cola = {5: " · armada falsa · SIN EMITIR", 6: " · Booking: FALSO778", 7: " · revisar falso", 10: " · error falso"}
        for fila, (estado, _) in estados.items():
            with self.subTest(fila=fila):
                texto = ws[f"H{fila}"].value
                self.assertIn(f" · {estado} · ", texto)
                self.assertTrue(texto.endswith(cola[fila]), texto)

    def test_estados_tras_el_envio_en_la_planilla(self):
        # Los estados de después de la guarda se escriben con su detalle y sin «SIN EMITIR»: esa
        # cola es solo de OK-EJEMPLO (decisión de Marcelo, CICLO-estado-tras-envio.md).
        estados = {5: ("NO ENVIADA", "ONE: no se pulsó «Submit» porque no estaba en la pantalla. La reserva no se envió."),
                   6: ("ENVIADA – REVISAR EN PORTAL", "ONE: se pulsó «Submit», pero la pantalla no mostró la "
                                                      "confirmación con el número de reserva."),
                   7: ("EMITIDA", "ONE emitida con éxito (Booking: SCLG10000004)"),
                   10: ("OK-EJEMPLO", "armada falsa")}
        self.poner_nave("ONE", 10, "NAVE PRUEBA TRECE")      # «NAVE POR COMPLETAR» ya no va al portal (CICLO-solo-la-nave-pedida.md)
        self.poner_nave("ONE", 7, "NAVE PRUEBA CATORCE")     # la columna F ya no es la nave (CICLO-cola-seis-items.md)
        with soporte.Navieras(self.mod, respuesta=lambda nav, rsv, on_pausa: estados[rsv["fila"]]):
            res, _ = self.mod.ejecutar_reservas("op_prueba", "one", esperar_cierre=lambda: None)
        self.assertEqual(res, {fila: e for fila, (e, _) in estados.items()})
        ws = openpyxl.load_workbook(self.planilla)["ONE"]
        for fila, (estado, detalle) in estados.items():
            with self.subTest(fila=fila):
                texto = ws[f"H{fila}"].value
                self.assertIn(f" · {estado} · nave: ", texto)
                self.assertTrue(texto.endswith(f" · {detalle}" + (" · SIN EMITIR" if fila == 10 else "")), texto)

    def poner_nave(self, hoja, fila, nave):
        wb = openpyxl.load_workbook(self.planilla)
        wb[hoja][f"E{fila}"] = nave
        wb.save(self.planilla)

    MOTIVO = "la fila no trae una nave para reservar"
    FECHA = r"\d\d-\d\d \d\d:\d\d"

    def linea_sin_nave(self, fila):
        return (f"· ✗ fila {fila}: NO ENVIADA · {self.MOTIVO}. No entro al portal por ella: escribe la nave en la "
                "planilla y vuelve a armar la reserva.")

    def test_fila_sin_nave_no_abre_el_portal(self):
        # Por la consola también: la fila sin nave queda NO ENVIADA con el motivo en la planilla y en log.txt, con su
        # número de solicitud, y no se abre el portal por ella (CICLO-solo-la-nave-pedida.md). En COSCO, la 5 queda sin nave y
        # la 6 con una: la que viene después conserva su número.
        # Primero, la hoja tal como viene (la 5 con nave, la 6 sin): la 6 se anota antes, pero el resultado y el resumen
        # van en el orden de la hoja.
        with soporte.Navieras(self.mod):
            res, carpeta = self.mod.ejecutar_reservas("op_prueba", "cosco", esperar_cierre=lambda: None)
        self.assertEqual(list(res.items()), [(5, "OK-EJEMPLO"), (6, "NO ENVIADA")])
        resumen = (Path(carpeta) / "log.txt").read_text(encoding="utf-8").split("===== RESUMEN RESERVAS =====")[1]
        self.assertLess(resumen.index("fila 5: OK-EJEMPLO"), resumen.index("fila 6: NO ENVIADA"))
        self.poner_nave("COSCO", 5, None)
        self.poner_nave("COSCO", 6, "NAVE PRUEBA TRECE")
        with soporte.Navieras(self.mod) as n:
            res, carpeta = self.mod.ejecutar_reservas("op_prueba", "cosco", esperar_cierre=lambda: None)
        self.assertEqual((res, [r[:2] for r in n.reservas]), ({5: "NO ENVIADA", 6: "OK-EJEMPLO"}, [("cosco", 6)]))
        self.assertEqual(list(res), [5, 6])
        ws = openpyxl.load_workbook(self.planilla)["COSCO"]
        self.assertRegex(ws["H5"].value, rf"^SOLICITUD 1 · NO ENVIADA · nave:  · PUERTO ALFA→CIUDAD GAMMA · {self.FECHA} "
                                         rf"· ref\(ej\) EJ\d{{6}}01 · {self.MOTIVO}$")
        self.assertRegex(ws["H6"].value, rf"^SOLICITUD 2 · OK-EJEMPLO · nave: NAVE PRUEBA TRECE · PUERTO ALFA→CIUDAD MU · "
                                         rf"{self.FECHA} · ref\(ej\) EJ\d{{6}}02 · armada falsa cosco fila 6 · SIN EMITIR$")
        log = (Path(carpeta) / "log.txt").read_text(encoding="utf-8")
        self.assertIn(self.linea_sin_nave(5), log)
        resumen = log.split("===== RESUMEN RESERVAS =====")[1]
        self.assertLess(resumen.index("fila 5: NO ENVIADA"), resumen.index("fila 6: OK-EJEMPLO"))
        # Con un comodín en la otra, ninguna trae nave: no se abre el navegador ni se entra al portal, y las dos quedan
        # en la planilla, con el resumen.
        self.poner_nave("COSCO", 6, "AUTO")
        with soporte.Navieras(self.mod) as n:
            res, carpeta = self.mod.ejecutar_reservas("op_prueba", "cosco", esperar_cierre=lambda: None)
        self.assertEqual((res, len(n.pw.lanzamientos), n.logins, n.reservas),
                         ({5: "NO ENVIADA", 6: "NO ENVIADA"}, 0, [], []))
        ws = openpyxl.load_workbook(self.planilla)["COSCO"]
        self.assertTrue(all(ws[f"H{f}"].value.endswith(f" · {self.MOTIVO}") for f in (5, 6)))
        log = (Path(carpeta) / "log.txt").read_text(encoding="utf-8")
        self.assertIn(self.linea_sin_nave(6), log)
        self.assertIn("· Ninguna fila de COSCO trae una nave para reservar: no abro el navegador ni entro al portal.", log)
        self.assertIn("· ===== RESUMEN RESERVAS =====", log)
        # Con solo el prefijo de la naviera, igual (CICLO-cola-seis-items.md): hasta f68ce6c iba al portal.
        self.poner_nave("COSCO", 6, "COSCO")
        with soporte.Navieras(self.mod) as n:
            res, _ = self.mod.ejecutar_reservas("op_prueba", "cosco", esperar_cierre=lambda: None)
        self.assertEqual((res, len(n.pw.lanzamientos), n.reservas), ({5: "NO ENVIADA", 6: "NO ENVIADA"}, 0, []))

    def test_fila_manual_queda_no_enviada(self):
        # En la consola no hay casillas: una fila MANUAL queda NO ENVIADA con su motivo en la planilla y en log.txt, sin
        # entrar al portal por ella (CICLO-cola-seis-items.md). Hasta f68ce6c iba al portal y buscaba «MANUAL» como nave.
        motivo = "la fila está marcada para hacerse a mano"
        self.poner_nave("COSCO", 6, "MANUAL")
        with soporte.Navieras(self.mod) as n:
            res, carpeta = self.mod.ejecutar_reservas("op_prueba", "cosco", esperar_cierre=lambda: None)
        self.assertEqual((res, [r[:2] for r in n.reservas]), ({5: "OK-EJEMPLO", 6: "NO ENVIADA"}, [("cosco", 5)]))
        ws = openpyxl.load_workbook(self.planilla)["COSCO"]
        self.assertRegex(ws["H6"].value, rf"^SOLICITUD 2 · NO ENVIADA · nave: MANUAL · .* · {motivo}$")
        self.assertIn(f"· ✗ fila 6: NO ENVIADA · {motivo}. No entro al portal por ella: hazla a mano en el portal, o "
                      "escribe en la celda solo la nave, sin MANUAL, si quieres que la arme yo.",
                      (Path(carpeta) / "log.txt").read_text(encoding="utf-8"))

    def test_la_nave_no_sale_de_la_columna_f(self):
        # En ONE y MSC, con la celda de la nave vacía, la consola ya no toma la nave de la columna F (en la planilla
        # estándar, «Viaje»): la fila queda sin nave, NO ENVIADA sin entrar al portal por ella, con el motivo en la
        # planilla y en log.txt (decisión de Marcelo, CICLO-cola-seis-items.md). Hasta f68ce6c iba al portal con el texto
        # de la F como nave.
        for nav, hoja, fila in (("one", "ONE", 7), ("msc", "MSC", 6)):
            with self.subTest(naviera=nav):
                with soporte.Navieras(self.mod) as n:
                    res, carpeta = self.mod.ejecutar_reservas("op_prueba", nav, esperar_cierre=lambda: None)
                self.assertEqual((res[fila], fila in [r[1] for r in n.reservas]), ("NO ENVIADA", False))
                texto = openpyxl.load_workbook(self.planilla)[hoja][f"H{fila}"].value
                self.assertRegex(texto, rf"^SOLICITUD \d · NO ENVIADA · nave:  · .* · {self.MOTIVO}$")
                self.assertIn(self.linea_sin_nave(fila), (Path(carpeta) / "log.txt").read_text(encoding="utf-8"))

    def test_fila_sin_nave_con_login_fallido_o_solo_la_primera(self):
        # La 5 sin nave y la 6 con una. Con el login fallido, la 5 queda igual en la planilla, con la copia de la
        # corrida, y la 6 no se toca: hasta a198645, un login fallido nunca tocaba la planilla
        # (test_login_fallido_no_toca_la_planilla, con filas que traen nave). Con AQUASHIELD_SOLO_PRIMERA, solo la 5, sin
        # abrir el navegador: la 6 no se toca (CICLO-solo-la-nave-pedida.md).
        self.poner_nave("COSCO", 5, None)
        self.poner_nave("COSCO", 6, "NAVE PRUEBA TRECE")
        antes = self.planilla.read_bytes()
        with soporte.Navieras(self.mod, login_ok=False) as n:
            res, _ = self.mod.ejecutar_reservas("op_prueba", "cosco", esperar_cierre=lambda: None)
        ws = openpyxl.load_workbook(self.planilla)["COSCO"]
        self.assertEqual((res, n.reservas, len(n.logins)), ({5: "NO ENVIADA"}, [], 1))
        self.assertEqual((ws["H5"].value.endswith(f" · {self.MOTIVO}"), ws["H6"].value), (True, None))
        self.assertEqual(len(list(self.sb.glob(RESPALDOS))), 1)
        self.planilla.write_bytes(antes)
        os.environ["AQUASHIELD_SOLO_PRIMERA"] = "1"
        try:
            with soporte.Navieras(self.mod) as n:
                res, _ = self.mod.ejecutar_reservas("op_prueba", "cosco", esperar_cierre=lambda: None)
        finally:
            os.environ.pop("AQUASHIELD_SOLO_PRIMERA", None)
        ws = openpyxl.load_workbook(self.planilla)["COSCO"]
        self.assertEqual((res, len(n.pw.lanzamientos), n.reservas, ws["H6"].value), ({5: "NO ENVIADA"}, 0, [], None))

    def poner(self, hoja, celdas):
        wb = openpyxl.load_workbook(self.planilla)
        for celda, valor in celdas.items():
            wb[hoja][celda] = valor
        wb.save(self.planilla)

    def test_fila_sin_puerto_o_sin_destino_no_abre_el_portal(self):
        # Por la consola también (leer_reservas): la fila con nave y sin puerto de carga, o sin destino, queda NO
        # ENVIADA con el motivo en la planilla y en log.txt, con su número de solicitud, sin abrir el portal por ella
        # (decisión de Marcelo, encargo 42, CICLO-maersk-y-fila-sin-ruta.md). En COSCO, la 5 sin puerto de carga y la 6
        # con nave.
        puerto, destino = "a la fila le falta el puerto de carga", "a la fila le falta el destino"
        self.poner("COSCO", {"B5": None, "E6": "NAVE PRUEBA TRECE"})
        with soporte.Navieras(self.mod) as n:
            res, carpeta = self.mod.ejecutar_reservas("op_prueba", "cosco", esperar_cierre=lambda: None)
        self.assertEqual((res, [r[:2] for r in n.reservas]), ({5: "NO ENVIADA", 6: "OK-EJEMPLO"}, [("cosco", 6)]))
        self.assertEqual(list(res), [5, 6])
        ws = openpyxl.load_workbook(self.planilla)["COSCO"]
        self.assertRegex(ws["H5"].value, rf"^SOLICITUD 1 · NO ENVIADA · nave: NAVE PRUEBA SIETE · →CIUDAD GAMMA · "
                                         rf"{self.FECHA} · ref\(ej\) EJ\d{{6}}01 · {puerto}$")
        log = (Path(carpeta) / "log.txt").read_text(encoding="utf-8")
        self.assertIn(f"· ✗ fila 5: NO ENVIADA · {puerto}. No entro al portal por ella: escribe el puerto de carga en "
                      "la planilla y vuelve a armar la reserva.", log)
        # Con la 6 sin destino, ninguna trae la ruta: no se abre el navegador ni se entra al portal, y las dos quedan en
        # la planilla, con el resumen.
        self.poner("COSCO", {"C6": None, "D6": None})
        with soporte.Navieras(self.mod) as n:
            res, carpeta = self.mod.ejecutar_reservas("op_prueba", "cosco", esperar_cierre=lambda: None)
        self.assertEqual((res, len(n.pw.lanzamientos), n.logins, n.reservas),
                         ({5: "NO ENVIADA", 6: "NO ENVIADA"}, 0, [], []))
        ws = openpyxl.load_workbook(self.planilla)["COSCO"]
        self.assertTrue(ws["H6"].value.endswith(f" · {destino}"), ws["H6"].value)
        log = (Path(carpeta) / "log.txt").read_text(encoding="utf-8")
        self.assertIn(f"· ✗ fila 6: NO ENVIADA · {destino}. No entro al portal por ella: escribe el destino en la "
                      "planilla y vuelve a armar la reserva.", log)
        self.assertIn("· Ninguna fila de COSCO con nave trae el puerto de carga y el destino: no abro el navegador ni "
                      "entro al portal.", log)
        self.assertIn("· ===== RESUMEN RESERVAS =====", log)

    def test_error_de_escritura_se_muestra_y_queda_en_el_log(self):
        # Planilla de solo lectura (como abierta en Excel): la corrida sigue, y cada resultado
        # que no se pudo escribir sale en pantalla (consola y ventana), en log.txt y en el
        # resumen, con el texto para anotarlo a mano. Hasta fba3bc3 se perdía en silencio.
        antes = self.planilla.read_bytes()
        vistas, salida = [], io.StringIO()
        os.chmod(self.planilla, stat.S_IREAD)
        try:
            with soporte.Navieras(self.mod), contextlib.redirect_stdout(salida):
                res, carpeta = self.mod.ejecutar_reservas("op_prueba", "cma", on_log=vistas.append,
                                                          esperar_cierre=lambda: None)
        finally:
            os.chmod(self.planilla, stat.S_IREAD | stat.S_IWRITE)
        self.assertEqual(res, {5: "OK-EJEMPLO", 6: "OK-EJEMPLO"})
        self.assertEqual(self.planilla.read_bytes(), antes)
        log = (Path(carpeta) / "log.txt").read_text(encoding="utf-8")
        motivo = ("«Reservas AQUASHIELD.xlsx» está abierta en otro programa (¿Excel?) o es de solo lectura; "
                  "ciérrala y vuelve a correr")
        avisos = [f"· ⚠ No pude escribir el resultado de la fila {fila} en la planilla: {motivo}. "
                  f"Anótalo a mano: SOLICITUD {i} · OK-EJEMPLO · nave: {nave} · "
                  for fila, i, nave in ((5, 1, "NAVE PRUEBA CINCO"), (6, 2, "NAVE PRUEBA SEIS"))]
        resumen = "· ⚠ 2 resultado(s) no quedaron en la planilla (filas 5, 6); cada uno está arriba en este log."
        for texto in avisos + [resumen]:
            with self.subTest(texto=texto[:60]):
                self.assertIn(texto, log)
                self.assertIn(texto, salida.getvalue())
                self.assertTrue(any(texto in v for v in vistas))

    def test_respaldo_antes_de_escribir(self):
        # Antes de la primera escritura, una copia con fecha y hora junto a la planilla: una
        # sola por corrida, con la planilla tal como estaba antes de escribir.
        antes = self.planilla.read_bytes()
        with soporte.Navieras(self.mod):
            _, carpeta = self.mod.ejecutar_reservas("op_prueba", "cma", esperar_cierre=lambda: None)
        respaldos = sorted(p.name for p in self.sb.glob(RESPALDOS))
        self.assertEqual(len(respaldos), 1, respaldos)
        self.assertRegex(respaldos[0], r"^Reservas AQUASHIELD - respaldo \d{4}-\d\d-\d\d \d{6}\.xlsx$")
        self.assertEqual((self.sb / respaldos[0]).read_bytes(), antes)
        self.assertNotEqual(self.planilla.read_bytes(), antes)
        self.assertIn(f"    Dejé una copia de respaldo de la planilla: {respaldos[0]}",
                      (Path(carpeta) / "log.txt").read_text(encoding="utf-8"))
        self.assertEqual(self.mod.encontrar_archivo_reservas(), self.planilla)   # la copia no pasa por planilla

    def test_todas_deja_una_sola_copia(self):
        # «todas» por consola corre las seis navieras seguidas y deja UNA copia, con la planilla de
        # antes de la primera escritura; las otras navieras la anotan en su log (hasta aaca5f1
        # dejaba una por naviera). Terminada la corrida, la siguiente hace la suya.
        antes = self.planilla.read_bytes()
        logs_antes = set((self.sb / "logs").glob("reservas_*"))
        sys.argv[:] = ["AQUASHIELD.py", "op_prueba", "todas", "reservas", "auto"]
        try:
            with soporte.Navieras(self.mod) as n:
                self.mod.main()
        finally:
            soporte.apagar_llaves()
        self.assertEqual(sorted({r[0] for r in n.reservas}), sorted(sinteticos.NAVIERAS))
        respaldos = sorted(self.sb.glob(RESPALDOS))
        self.assertEqual(len(respaldos), 1, [p.name for p in respaldos])
        self.assertEqual(respaldos[0].read_bytes(), antes)
        textos = [(c / "log.txt").read_text(encoding="utf-8")
                  for c in set((self.sb / "logs").glob("reservas_*")) - logs_antes]
        nombre = respaldos[0].name
        self.assertEqual((len(textos),
                          sum(f"Dejé una copia de respaldo de la planilla: {nombre}" in t for t in textos),
                          sum(f"La copia de respaldo de esta corrida ya está: {nombre}" in t for t in textos)),
                         (6, 1, 5))
        with soporte.Navieras(self.mod):
            self.mod.ejecutar_reservas("op_prueba", "cma", esperar_cierre=lambda: None)
        self.assertEqual(len(list(self.sb.glob(RESPALDOS))), 2)

    def test_poda_deja_las_10_mas_recientes_y_nada_mas(self):
        # Solo archivos de ESA carpeta cuyo nombre calce exacto. La antigüedad sale del nombre: las
        # fechas de archivo van al revés a propósito. Los señuelos son más viejos que todas las copias,
        # así que cualquier coincidencia de más los borraría primero.
        carpeta = self.sb / "poda"
        shutil.rmtree(carpeta, ignore_errors=True)
        (carpeta / "sub").mkdir(parents=True)
        exactos = [f"Reservas AQUASHIELD - respaldo 2026-01-{d:02d} 120000.xlsx" for d in range(1, 13)]
        senuelos = ["Reservas AQUASHIELD - respaldo 2025-01-01 000000 (2).xlsx",
                    "reservas aquashield - respaldo 2025-01-01 000001.xlsx",
                    "Reservas AQUASHIELD - respaldo 2025-01-01 000002.xlsm",
                    "Reservas AQUASHIELD - respaldo 2025-01-01 000003.xlsx.bak",
                    "Agendamientos - respaldo 2025-01-01 000004.xlsx",
                    "Copia de Reservas AQUASHIELD - respaldo 2025-01-01 000005.xlsx",
                    "Reservas AQUASHIELD - respaldo 2025-1-1 000006.xlsx",
                    "Reservas AQUASHIELD.xlsx", "config.json",
                    "sub/Reservas AQUASHIELD - respaldo 2025-01-01 000007.xlsx"]
        carpeta_exacta = "Reservas AQUASHIELD - respaldo 2025-01-01 000008.xlsx"
        for i, nombre in enumerate(exactos):
            (carpeta / nombre).write_bytes(b"copia sintetica")
            os.utime(carpeta / nombre, (2_000_000_000 - i * 1000,) * 2)
        for nombre in senuelos:
            (carpeta / nombre).write_bytes(b"no es una copia")
        (carpeta / carpeta_exacta).mkdir()
        borrados, errores = self.mod.podar_respaldos(carpeta)
        self.assertEqual((borrados, errores), (exactos[:2], []))
        quedan = {p.relative_to(carpeta).as_posix() for p in carpeta.rglob("*")}
        self.assertEqual(quedan, set(exactos[2:]) | set(senuelos) | {"sub", carpeta_exacta})
        self.assertEqual(self.mod.podar_respaldos(carpeta), ([], []))     # con 10, no borra nada

    def test_respaldo_poda_las_antiguas(self):
        # Al dejar la copia de la corrida quedan las 10 más recientes. La que no se puede borrar (de
        # solo lectura) queda y se avisa, y la corrida sigue escribiendo.
        viejas = [self.sb / f"Reservas AQUASHIELD - respaldo 2025-01-{d:02d} 120000.xlsx" for d in range(1, 12)]
        for p in viejas:
            p.write_bytes(b"copia sintetica vieja")
        os.chmod(viejas[0], stat.S_IREAD)
        antes = self.planilla.read_bytes()
        try:
            with soporte.Navieras(self.mod):
                _, carpeta = self.mod.ejecutar_reservas("op_prueba", "cma", esperar_cierre=lambda: None)
        finally:
            os.chmod(viejas[0], stat.S_IREAD | stat.S_IWRITE)
        nombres = sorted(p.name for p in self.sb.glob(RESPALDOS))
        nueva = [n for n in nombres if n not in {p.name for p in viejas}]
        self.assertEqual(len(nueva), 1, nombres)
        self.assertEqual(nombres, sorted([viejas[0].name] + [p.name for p in viejas[2:]] + nueva))
        log = (Path(carpeta) / "log.txt").read_text(encoding="utf-8")
        self.assertIn("    Borré 1 copia(s) de respaldo antiguas (se conservan las 10 más recientes).", log)
        self.assertIn(f"· ⚠ No pude borrar la copia de respaldo antigua «{viejas[0].name}» (PermissionError: ", log)
        self.assertNotEqual(self.planilla.read_bytes(), antes)

    def test_si_la_poda_falla_igual_escribe(self):
        antes = self.planilla.read_bytes()
        orig = self.mod.podar_respaldos

        def falla(carpeta):
            raise OSError("carpeta ilegible de prueba")
        self.mod.podar_respaldos = falla
        try:
            with soporte.Navieras(self.mod):
                _, carpeta = self.mod.ejecutar_reservas("op_prueba", "cma", esperar_cierre=lambda: None)
        finally:
            self.mod.podar_respaldos = orig
        self.assertEqual(len(list(self.sb.glob(RESPALDOS))), 1)
        self.assertNotEqual(self.planilla.read_bytes(), antes)
        self.assertIn("· ⚠ No pude revisar las copias de respaldo antiguas (OSError: carpeta ilegible de prueba); "
                      "la corrida sigue.", (Path(carpeta) / "log.txt").read_text(encoding="utf-8"))

    def test_sin_respaldo_no_escribe(self):
        # Si la copia falla, no se escribe encima: cada fila avisa y el resumen las cuenta.
        antes = self.planilla.read_bytes()

        def falla(*a, **k):
            raise PermissionError(13, "sin permiso de prueba")
        orig = shutil.copyfile
        shutil.copyfile = falla
        try:
            with soporte.Navieras(self.mod):
                res, carpeta = self.mod.ejecutar_reservas("op_prueba", "cma", esperar_cierre=lambda: None)
        finally:
            shutil.copyfile = orig
        self.assertEqual(res, {5: "OK-EJEMPLO", 6: "OK-EJEMPLO"})
        self.assertEqual(self.planilla.read_bytes(), antes)
        self.assertEqual(list(self.sb.glob(RESPALDOS)), [])
        log = (Path(carpeta) / "log.txt").read_text(encoding="utf-8")
        self.assertIn("· ⚠ No pude escribir el resultado de la fila 5 en la planilla: no pude dejar la copia de "
                      "respaldo de «Reservas AQUASHIELD.xlsx» (PermissionError: [Errno 13] sin permiso de prueba); "
                      "sin respaldo no escribo encima. Anótalo a mano: SOLICITUD 1 · OK-EJEMPLO · ", log)
        self.assertIn("(filas 5, 6)", log)

    def test_respaldos_no_se_pisan(self):
        # Tres respaldos seguidos (casi siempre en el mismo segundo) dejan tres archivos.
        rutas = [self.mod.respaldar_planilla() for _ in range(3)]
        self.assertEqual(len({r.name for r in rutas}), 3)
        self.assertEqual(sorted(p.name for p in self.sb.glob(RESPALDOS)), sorted(r.name for r in rutas))

    def test_sin_escrituras_no_deja_respaldo(self):
        # Con HYUNDAI, que no deja la fila NO ENVIADA si su login falla; MSC sí, desde el encargo 45
        # (test_login_fallido_de_msc_queda_no_enviada).
        with soporte.Navieras(self.mod, login_ok=False):
            self.mod.ejecutar_reservas("op_prueba", "hyundai", esperar_cierre=lambda: None)
        self.assertEqual(list(self.sb.glob(RESPALDOS)), [])

    def test_solo_primera_fila(self):
        os.environ["AQUASHIELD_SOLO_PRIMERA"] = "1"
        try:
            with soporte.Navieras(self.mod) as n:
                res, _ = self.mod.ejecutar_reservas("op_prueba", "hyundai", esperar_cierre=lambda: None)
        finally:
            os.environ.pop("AQUASHIELD_SOLO_PRIMERA", None)
        self.assertEqual((res, len(n.reservas)), ({5: "OK-EJEMPLO"}, 1))

    def test_login_fallido_no_toca_la_planilla(self):
        # Con HYUNDAI: el login fallido de las navieras que no son MSC no toca la planilla (encargo 45; hasta ahí, MSC).
        antes = self.planilla.read_bytes()
        with soporte.Navieras(self.mod, login_ok=False) as n:
            res, _ = self.mod.ejecutar_reservas("op_prueba", "hyundai", esperar_cierre=lambda: None)
        self.assertEqual((res, n.reservas), ({}, []))
        self.assertEqual(self.planilla.read_bytes(), antes)

    def test_login_fallido_de_msc_queda_no_enviada(self):
        # Sin la sesión de MSC, cada fila con nave queda NO ENVIADA con su motivo en la planilla, con la copia de la
        # corrida antes (decisión de Marcelo, encargo 45, CICLO-login-msc-y-maersk.md); hasta ahí, no se tocaba.
        self.poner_nave("MSC", 6, "NAVE PRUEBA QUINCE")      # la columna F ya no es la nave (CICLO-cola-seis-items.md)
        with soporte.Navieras(self.mod, login_ok=False) as n:
            res, _ = self.mod.ejecutar_reservas("op_prueba", "msc", esperar_cierre=lambda: None)
        # La 8 es el encabezado repetido de la hoja sintética: el lector de la consola la lee como fila, como antes
        # (test_planilla, CONSOLA["msc"]).
        self.assertEqual((res, n.reservas), ({5: "NO ENVIADA", 6: "NO ENVIADA", 7: "NO ENVIADA", 8: "NO ENVIADA"}, []))
        ws = openpyxl.load_workbook(self.planilla)["MSC"]
        for f in (5, 6, 7, 8):
            with self.subTest(fila=f):
                self.assertIn(" · NO ENVIADA · ", ws[f"H{f}"].value)
                self.assertTrue(ws[f"H{f}"].value.endswith(self.mod.MSC_SIN_SESION))
        self.assertEqual(len(list(self.sb.glob(RESPALDOS))), 1)

    def test_login_de_msc_deja_su_motivo(self):
        # Si el login de MSC dejó su motivo en el Registro (tras el segundo intento: MSC_SIN_SESION_TRAS_DOS,
        # decisión de Marcelo, encargo 50), es el de cada fila de MSC en la planilla (_motivo_sin_sesion).
        motivo = self.mod._msc_motivo_tras_dos(self.mod.MSC_ERROR_TRAS_LA_CLAVE, "")

        def login(page, creds, reg, on_pausa=None):
            reg.sin_sesion = {"msc": motivo}
            return False
        with soporte.Navieras(self.mod, login_ok=False) as n:
            self.mod.NAVIERAS["msc"] = ("MSC", login)
            res, _ = self.mod.ejecutar_reservas("op_prueba", "msc", esperar_cierre=lambda: None)
        self.assertEqual((res.get(5), n.reservas), ("NO ENVIADA", []))
        self.assertTrue(openpyxl.load_workbook(self.planilla)["MSC"]["H5"].value.endswith(motivo))

    def test_cma_detenida_en_el_login_queda_no_enviada(self):
        # Si el login de CMA-CGM se detuvo porque el portal restringió el acceso o no terminó de cargar (decisión de
        # Marcelo, encargo 54), cada fila de CMA queda NO ENVIADA con ese motivo en la planilla, sin reservar; hasta ahí
        # no se tocaba, como con el login fallido de las navieras que no son MSC.
        motivo = self.mod.CMA_DETENIDA.format(causa=self.mod.CMA_ACCESO_RESTRINGIDO)

        def login(page, creds, reg, on_pausa=None):
            self.mod._cma_detener(page, reg, self.mod.CMA_ACCESO_RESTRINGIDO, "cma_acceso_restringido", creds)
        with soporte.Navieras(self.mod) as n:
            self.mod.NAVIERAS["cma"] = ("CMA-CGM", self.mod._cma_login_se_detiene(login))
            res, _ = self.mod.ejecutar_reservas("op_prueba", "cma", esperar_cierre=lambda: None)
        self.assertEqual((sorted(res), sorted(set(res.values())), n.reservas), ([5, 6], ["NO ENVIADA"], []))
        ws = openpyxl.load_workbook(self.planilla)["CMA-CGM"]
        for f in res:
            with self.subTest(fila=f):
                self.assertTrue(ws[f"H{f}"].value.endswith(motivo))

    def test_reservador_que_revienta(self):
        def respuesta(nav, rsv, on_pausa):
            raise RuntimeError("falla falsa")
        with soporte.Navieras(self.mod, respuesta=respuesta) as n:
            res, _ = self.mod.ejecutar_reservas("op_prueba", "maersk", esperar_cierre=lambda: None)
        self.assertEqual(res, {5: "ERROR", 6: "ERROR"})
        self.assertEqual(n.pw.contextos[0].pages[0].capturas, [("error_fila_5.png", False), ("error_fila_6.png", False)])
        self.assertIn("· ERROR · ", openpyxl.load_workbook(self.planilla)["MAERSK"]["H5"].value)

    def test_dice_su_candado_al_empezar(self):
        # Cada corrida de reservas por consola dice al empezar, en log.txt, con qué candado corre y por cuál llave
        # (_anotar_candado; decisión de Marcelo, encargo 51). Con llaves, aquí dichas por _llaves_abiertas solo para la
        # línea: el candado de la guarda sigue cerrado.
        with soporte.Navieras(self.mod):
            _, carpeta = self.mod.ejecutar_reservas("op_prueba", "cma", esperar_cierre=lambda: None)
        lineas = (Path(carpeta) / "log.txt").read_text(encoding="utf-8").splitlines()
        self.assertIn("· INICIO RESERVAS · CMA-CGM · usuario=op_prueba", lineas[0])
        self.assertIn("· 🛡️ Candado de emisión cerrado: modo prueba", lineas[1])
        llaves = (lambda todas=False: ["la llave de prueba"] if todas else [])
        with soporte.Navieras(self.mod), mock.patch.object(self.mod, "_llaves_abiertas", llaves):
            _, carpeta = self.mod.ejecutar_reservas("op_prueba", "cma", esperar_cierre=lambda: None)
        lineas = (Path(carpeta) / "log.txt").read_text(encoding="utf-8").splitlines()
        self.assertIn("· 🔴 Candado de emisión abierto: modo EMISIÓN, por la llave de prueba.", lineas[1])

    def test_rechazos(self):
        with self.assertRaises(ValueError) as e:
            self.mod.ejecutar_reservas("nadie", "one", esperar_cierre=lambda: None)
        self.assertIn("no está en config.json", str(e.exception))
        with self.assertRaises(ValueError) as e:
            self.mod.ejecutar_reservas("op_prueba", "xx", esperar_cierre=lambda: None)
        self.assertIn("Aún no está lista la reserva automática de 'xx'", str(e.exception))


class TestEjecutarLogin(ConPlanilla):
    def test_login_por_consola(self):
        with soporte.Navieras(self.mod, login_ok=lambda nav: nav != "msc") as n:
            res, carpeta = self.mod.ejecutar_login("op_prueba", ["one", "msc", "xx"], esperar_cierre=lambda: None)
        self.assertEqual(res, {"one": True, "msc": False})
        self.assertEqual([l[0] for l in n.logins], ["one", "msc"])
        self.assertRegex(Path(carpeta).name, r"^op_prueba_\d{8}_\d{6}$")
        self.assertEqual(Path(n.pw.lanzamientos[0]["user_data_dir"]), self.sb / "perfiles" / "op_prueba")
        self.assertIs(n.pw.lanzamientos[0]["chromium_sandbox"], True)      # con el sandbox de Chrome (encargo 53)
        self.assertTrue(n.pw.contextos[0].cerrado)

    def test_sin_credenciales_igual_abre_el_navegador(self):
        with soporte.Navieras(self.mod) as n:
            res, _ = self.mod.ejecutar_login("op_vacio", ["one", "msc"], esperar_cierre=lambda: None)
        self.assertEqual(res, {"one": None, "msc": None})
        self.assertEqual((n.logins, len(n.pw.lanzamientos)), ([], 1))

    def test_login_que_revienta(self):
        with soporte.Navieras(self.mod) as n:
            def rompe(page, creds, reg, on_pausa=None):
                raise RuntimeError("login falso roto")
            self.mod.NAVIERAS["cosco"] = ("COSCO", rompe)
            res, _ = self.mod.ejecutar_login("op_prueba", ["cosco"], esperar_cierre=lambda: None)
        self.assertEqual(res, {"cosco": False})
        self.assertEqual(n.pw.contextos[0].pages[0].capturas, [("cosco_error.png", False)])

    def test_login_que_se_detiene_en_cma_sigue(self):
        # En «Solo iniciar sesión», CMA-CGM detenida por el portal queda REVISAR, con su línea en log.txt, y la naviera
        # siguiente entra igual: la detención no corta la corrida (decisión de Marcelo, encargo 54). Corre después de
        # test_login_por_consola, que pide su carpeta sin «_2».
        def login(page, creds, reg, on_pausa=None):
            self.mod._cma_detener(page, reg, self.mod.CMA_ACCESO_RESTRINGIDO, "cma_acceso_restringido", creds)
        with soporte.Navieras(self.mod) as n:
            self.mod.NAVIERAS["cma"] = ("CMA-CGM", self.mod._cma_login_se_detiene(login))
            res, carpeta = self.mod.ejecutar_login("op_prueba", ["cma", "cosco"], esperar_cierre=lambda: None)
        self.assertEqual((res, [l[0] for l in n.logins]), ({"cma": False, "cosco": True}, ["cosco"]))
        log = (Path(carpeta) / "log.txt").read_text(encoding="utf-8")
        self.assertIn("· ⛔ CMA-CGM restringió el acceso: su página dice «El acceso está restringido temporalmente». "
                      "Me detengo sin reintentar ni recargar la página.", log)
        self.assertIn("· [CMA-CGM] resultado=REVISAR", log)

    def test_operador_desconocido(self):
        with self.assertRaises(ValueError):
            self.mod.ejecutar_login("nadie", ["one"], esperar_cierre=lambda: None)

    def test_dice_su_candado_al_empezar(self):
        # También la corrida de solo login dice al empezar con qué candado corre (encargo 51): no emite, pero es el
        # mismo programa, en el mismo modo, que arma las reservas. Con el operador sin credenciales: su carpeta no
        # choca en el mismo segundo con la de test_login_por_consola, que pide el nombre sin «_2».
        with soporte.Navieras(self.mod):
            _, carpeta = self.mod.ejecutar_login("op_vacio", ["one"], esperar_cierre=lambda: None)
        lineas = (Path(carpeta) / "log.txt").read_text(encoding="utf-8").splitlines()
        self.assertIn("· INICIO · usuario=op_vacio · navieras=one", lineas[0])
        self.assertIn("· 🛡️ Candado de emisión cerrado: modo prueba", lineas[1])

    def test_solo_login_con_detener_no_sigue(self):
        # Decisión de Marcelo, encargo 57 (CICLO-cortes-de-cma-y-pausas.md): con «Detener» (se_detuvo), no sigue con los
        # inicios de sesión que faltan, no espera para cerrar y cierra el navegador; al_abrir recibe el navegador apenas
        # se abre (el panel web lo guarda para que «Detener» lo cierre). Aquí el operador lo pulsa durante el de ONE. Si
        # el inicio de sesión se cae porque «Detener» le cerró el navegador, no es un error del portal: no hay captura.
        # Sin se_detuvo (la consola y el panel Tkinter), como antes. Corre después de test_login_por_consola, que pide
        # su carpeta sin «_2».
        pulsado, abiertos, esperas = [], [], []

        def rompe(page, creds, reg, on_pausa=None):
            pulsado.append("one")
            raise RuntimeError("Target page, context or browser has been closed")
        with soporte.Navieras(self.mod) as n:
            self.mod.NAVIERAS["one"] = ("ONE", rompe)
            res, carpeta = self.mod.ejecutar_login("op_prueba", ["one", "msc", "cosco"],
                                                   esperar_cierre=lambda: esperas.append(1),
                                                   se_detuvo=lambda: bool(pulsado), al_abrir=abiertos.append)
        self.assertEqual((res, n.logins, esperas), ({"one": False}, [], []))
        self.assertEqual((abiertos, n.pw.contextos[0].cerrado, n.pw.contextos[0].pages[0].capturas),
                         ([n.pw.contextos[0]], True, []))
        log = (Path(carpeta) / "log.txt").read_text(encoding="utf-8")
        self.assertIn("    se cortó con «Detener» (RuntimeError: Target page, context or browser has been closed)", log)
        self.assertIn("· ⛔ Detenido por el operador: cierro el navegador y no sigo con MSC, COSCO.", log)
        self.assertNotIn("ERROR inesperado", log)
        # Pulsado después del último, lo dice sin la lista de los que faltan, y tampoco espera.
        with soporte.Navieras(self.mod) as n:
            res, carpeta = self.mod.ejecutar_login("op_prueba", ["one"], esperar_cierre=lambda: esperas.append(1),
                                                   se_detuvo=lambda: bool(n.logins))
        log = (Path(carpeta) / "log.txt").read_text(encoding="utf-8")
        self.assertEqual((res, esperas), ({"one": True}, []))
        self.assertIn("· ⛔ Detenido por el operador: cierro el navegador.\n", log)


class TestCarpetaCorrida(ConPlanilla):
    # Hasta 393f504, dos corridas del mismo segundo compartían carpeta: la segunda escribía en el log de la
    # primera y pisaba sus capturas (CICLO-carpetas-unicas.md). El reloj detenido las hace chocar siempre.
    def inicios(self, carpeta):
        return len(re.findall(r"· INICIO", (Path(carpeta) / "log.txt").read_text(encoding="utf-8")))

    def test_mismo_segundo_da_carpetas_distintas(self):
        with soporte.reloj_detenido(self.mod):
            carpetas = [self.mod.carpeta_corrida("prueba_choque") for _ in range(3)]
            otra = self.mod.carpeta_corrida("prueba_otra")
        self.assertEqual([c.name for c in carpetas], ["prueba_choque_20260102_030405", "prueba_choque_20260102_030405_2",
                                                      "prueba_choque_20260102_030405_3"])
        self.assertEqual(otra.name, "prueba_otra_20260102_030405")     # sin choque, el nombre de siempre
        self.assertTrue(all(c.is_dir() and c.parent == self.sb / "logs" for c in carpetas + [otra]))
        # El orden alfabético sigue siendo el de las corridas.
        self.assertEqual(sorted(c.name for c in carpetas), [c.name for c in carpetas])

    def test_consola_en_el_mismo_segundo(self):
        with soporte.reloj_detenido(self.mod), soporte.Navieras(self.mod):
            carpetas = [self.mod.ejecutar_reservas("op_prueba", "cma", esperar_cierre=lambda: None)[1]
                        for _ in range(2)]
        self.assertEqual([Path(c).name for c in carpetas],
                         ["reservas_cma_op_prueba_20260102_030405", "reservas_cma_op_prueba_20260102_030405_2"])
        self.assertEqual([self.inicios(c) for c in carpetas], [1, 1])

    def test_login_en_el_mismo_segundo(self):
        with soporte.reloj_detenido(self.mod), soporte.Navieras(self.mod):
            carpetas = [self.mod.ejecutar_login("op_prueba", ["one"], esperar_cierre=lambda: None)[1]
                        for _ in range(2)]
        self.assertEqual([Path(c).name for c in carpetas], ["op_prueba_20260102_030405", "op_prueba_20260102_030405_2"])
        self.assertEqual([self.inicios(c) for c in carpetas], [1, 1])


class TestUtilitarios(soporte.CasoAQ):
    def test_registro(self):
        vistas = []
        reg = self.mod.Registro(self.sb / "reg" / "sub" / "log.txt", on_log=vistas.append)
        reg.paso("hola")
        reg.info("dato")
        reg.cerrar()
        lineas = (self.sb / "reg" / "sub" / "log.txt").read_text(encoding="utf-8").splitlines()
        self.assertEqual(lineas, vistas)
        self.assertRegex(lineas[0], r"^\[\d\d:\d\d:\d\d \+\s+\d+\.\ds\] · hola$")
        self.assertRegex(lineas[1], r"^\[\d\d:\d\d:\d\d \+\s+\d+\.\ds\]     dato$")

    def test_url_sin_consulta(self):
        # La línea «URL:» de log.txt va sin la consulta (lo que va desde «?») ni el fragmento (desde «#»), y sin un
        # usuario o una clave antes del servidor, en las seis navieras, y dice qué quitó (Registro.url con
        # _sin_consulta; decisión de Marcelo, encargo 51). Medido en logs/: 40 líneas traían el usuario de COSCO
        # (login_hint) y 18, el estado o el código del inicio de sesión de ONE. Direcciones y datos inventados.
        casos = [
            ("https://iam.ejemplo.test/auth/realms/r/protocol/openid-connect/auth?client_id=c&login_hint="
             "usuario%40ejemplo.test&ui_locales=es",
             "https://iam.ejemplo.test/auth/realms/r/protocol/openid-connect/auth (sin su consulta)"),
            ("https://www.ejemplo.test/ecom/booking?code=CODIGODEPRUEBA&state=ESTADODEPRUEBA",
             "https://www.ejemplo.test/ecom/booking (sin su consulta)"),
            ("https://app.ejemplo.test/inicio#access_token=TOKENDEPRUEBA",
             "https://app.ejemplo.test/inicio (sin su fragmento)"),
            ("https://usuario:CLAVEDEPRUEBA@portal.ejemplo.test:8443/a/b?x=1#y",
             "https://portal.ejemplo.test:8443/a/b (sin su consulta ni su fragmento)"),
            ("https://www.ejemplo.test/Ruta/Sin/Nada", "https://www.ejemplo.test/Ruta/Sin/Nada"),
            ("about:blank", "about:blank"),
            ("chrome-error://chromewebdata/", "chrome-error://chromewebdata/"),
            ("http://[::1]:8443/x?y=1", "http://[::1]:8443/x (sin su consulta)"),       # IPv6, con sus corchetes
        ]

        class SinDireccion:
            @property
            def url(self):
                raise RuntimeError("página cerrada")
        reg = self.mod.Registro(self.sb / "url" / "log.txt")
        try:
            for url, _ in casos:
                reg.url(soporte.PaginaFalsa(url=url))
            reg.url(SinDireccion())
        finally:
            reg.cerrar()
        texto = (self.sb / "url" / "log.txt").read_text(encoding="utf-8")
        self.assertEqual(re.findall(r"\]     URL: (.*)", texto), [e for _, e in casos])
        for secreto in ("usuario%40", "login_hint", "CODIGODEPRUEBA", "ESTADODEPRUEBA", "TOKENDEPRUEBA",
                        "CLAVEDEPRUEBA", "usuario:"):
            self.assertNotIn(secreto, texto)
        self.assertEqual(self.mod._sin_consulta("https://x.ejemplo.test:puerto/"), "(no pude leer la dirección)")
        # El avance de MAERSK, igual: la línea «avance detectado por URL:» va sin la consulta.
        reg = self.mod.Registro(self.sb / "url_mk" / "log.txt")
        try:
            pagina = soporte.PaginaFalsa(url="https://www.maersk.ejemplo.test/book/sailings?token=TOKENDEPRUEBA")
            with mock.patch.object(self.mod, "esperar", lambda *a, **k: None):
                self.assertIs(self.mod._mk_esperar_avance(pagina, "", 5, reg), True)
        finally:
            reg.cerrar()
        avance = (self.sb / "url_mk" / "log.txt").read_text(encoding="utf-8")
        self.assertIn("avance detectado por URL: https://www.maersk.ejemplo.test/book/sailings (sin su consulta) (",
                      avance)
        self.assertNotIn("tokendeprueba", avance.lower())

    def test_one_paso_de_la_url(self):
        # Lo que escribe ONE de su dirección en log.txt y en el motivo (_one_paso_de_la_url): el valor de «step=», hasta
        # el «&» siguiente, o la dirección sin su consulta si no lo trae; nunca otro parámetro (revisión del encargo 51:
        # reservar_one escribía la dirección entera si no traía «step=»). Direcciones inventadas.
        f = self.mod._one_paso_de_la_url
        casos = {"https://www.one.ejemplo.test/ecom/booking?step=booking-parties": "booking-parties",
                 "https://www.one.ejemplo.test/ecom/booking?step=review-booking&code=CODIGODEPRUEBA": "review-booking",
                 "https://www.one.ejemplo.test/ecom/?code=CODIGODEPRUEBA&state=ESTADODEPRUEBA":
                     "https://www.one.ejemplo.test/ecom/",
                 "https://www.one.ejemplo.test/ecom/booking?step=" + "x" * 40: "x" * 30, "": "", None: ""}
        for url, esperado in casos.items():
            with self.subTest(url=url):
                self.assertEqual(f(url), esperado)

    def test_ninguna_direccion_con_su_consulta_a_un_mensaje(self):
        # Ninguna dirección de una página llega con su consulta a un mensaje: ni a una línea de log.txt (reg.paso,
        # reg.info, _emit) ni al motivo que devuelve una reserva (revisión del encargo 51: reservar_one escribía la
        # dirección entera). Por función: los nombres que toman algo de una dirección (.url), salvo un sí o un no, y
        # salvo que pase por _sin_consulta o _one_paso_de_la_url; ni ellos ni una dirección leída ahí mismo pueden
        # llegar a un mensaje sin pasar por esas dos. Control positivo: una función inventada que lo hace, cae.
        limpiadores = ("_sin_consulta(", "_one_paso_de_la_url(")

        def sucias(arbol):
            out = []
            for f in ast.walk(arbol):
                if not isinstance(f, ast.FunctionDef):
                    continue
                nombres = set()
                for n in ast.walk(f):
                    if isinstance(n, ast.Assign):
                        v = ast.unparse(n.value)
                        booleano = isinstance(n.value, (ast.Compare, ast.BoolOp, ast.UnaryOp)) or (
                            isinstance(n.value, ast.Call) and ast.unparse(n.value.func) in ("any", "all", "bool"))
                        if ".url" in v and not booleano and not any(x in v for x in limpiadores):
                            nombres |= {t.id for t in n.targets if isinstance(t, ast.Name)}
                for n in ast.walk(f):
                    es_log = (isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute)
                              and n.func.attr in ("paso", "info", "_emit"))
                    es_motivo = isinstance(n, ast.Return) and isinstance(n.value, ast.Tuple)
                    if not (es_log or es_motivo):
                        continue
                    texto = ast.unparse(n)
                    usa = ({x.id for x in ast.walk(n) if isinstance(x, ast.Name)} & nombres) or ".url" in texto
                    if usa and not any(x in texto for x in limpiadores):
                        out.append(f"{f.name}: {texto[:90]}")
            return out
        control = ast.parse("def f(page, reg):\n    a = page.url.split('x')[0]\n    reg.paso(f'quedó en {a}')\n"
                            "    return ('REVISAR', f'en {page.url}')\n")
        self.assertEqual(len(sucias(control)), 2)
        self.assertEqual(sucias(ast.parse(Path(self.mod.__file__).read_text(encoding="utf-8"))), [])

    def test_toda_linea_url_va_sin_consulta(self):
        # Ninguna línea con «URL:» escribe una dirección sin pasar por _sin_consulta (encargo 51). Hoy son dos:
        # Registro.url, que usan las seis navieras, y el avance de MAERSK. Una nueva que escriba page.url entera, cae.
        tree = ast.parse(Path(self.mod.__file__).read_text(encoding="utf-8"))
        lineas = {}
        for f in ast.walk(tree):
            if isinstance(f, ast.FunctionDef):
                for c in ast.walk(f):
                    if (isinstance(c, ast.Call) and isinstance(c.func, ast.Attribute)
                            and c.func.attr in ("paso", "info", "_emit") and "URL:" in ast.unparse(c)):
                        lineas[f.name] = ast.unparse(c)
        self.assertEqual(sorted(lineas), ["_mk_esperar_avance", "url"])
        self.assertEqual({f: "_sin_consulta(" in c for f, c in lineas.items()}, {f: True for f in lineas})

    def test_capturas(self):
        reg = self.mod.Registro(self.sb / "cap" / "log.txt")
        pagina = soporte.PaginaFalsa()
        for env in ("AQUASHIELD_CAPTURA_FULL", "AQUASHIELD_DESCUBRIR"):
            os.environ.pop(env, None)
        reg.captura(pagina, "a")
        reg.captura(pagina, "b", full=True)                    # sin variable: igual solo lo visible
        for env in ("AQUASHIELD_CAPTURA_FULL", "AQUASHIELD_DESCUBRIR"):
            os.environ[env] = "1"
            try:
                reg.captura(pagina, f"c_{env[-4:].lower()}", full=True)
            finally:
                os.environ.pop(env, None)
        self.assertEqual(pagina.capturas, [("a.png", False), ("b.png", False), ("c_full.png", True), ("c_brir.png", True)])

        class Rompe(soporte.PaginaFalsa):
            def screenshot(self, path=None, full_page=False):
                raise RuntimeError("sin pantalla")
        reg.captura(Rompe(), "d")
        reg.cerrar()
        self.assertIn("(no pude capturar d: sin pantalla)", (self.sb / "cap" / "log.txt").read_text(encoding="utf-8"))

    def test_pausa_manual(self):
        reg = self.mod.Registro(self.sb / "pausa" / "log.txt")
        avisos = []
        self.mod.pausa_manual(reg, avisos.append, "resuelve esto")
        self.assertEqual(avisos, ["resuelve esto"])
        # Sin panel ni consola: espera AQUASHIELD_PAUSA_SEG (25 s por defecto) y sigue.
        esperas = []
        orig_time, orig_stdin = self.mod.time, sys.stdin
        self.mod.time = types.SimpleNamespace(sleep=esperas.append, time=orig_time.time)
        sys.stdin = io.StringIO()
        os.environ.pop("AQUASHIELD_PAUSA_SEG", None)
        try:
            self.mod.pausa_manual(reg, None, "otra")
            os.environ["AQUASHIELD_PAUSA_SEG"] = "3"
            self.mod.pausa_manual(reg, None, "otra")
        finally:
            self.mod.time, sys.stdin = orig_time, orig_stdin
            os.environ.pop("AQUASHIELD_PAUSA_SEG", None)
        self.assertEqual(esperas, [25, 3])
        reg.cerrar()

    def test_pausa_del_panel_en_log(self):
        # Decisión de Marcelo, encargo 55 (CICLO-pausa-plazos-y-datadome.md): la pausa del panel queda en log.txt,
        # cuándo empieza (la hora de la línea) y por qué, y cuándo termina y cómo, o si se cortó. Hasta ahí solo la
        # veía el registro del panel: el 2026-10-06, log.txt quedó 6 minutos sin escribir.
        m = self.mod
        reloj = types.SimpleNamespace(t=100.0)
        reloj.time = lambda: reloj.t
        casos = [(m.PAUSA_RESUELTA, 12, "· ▶ La pausa terminó a los 12 s: el operador pulsó «Ya lo resolví»."),
                 (m.PAUSA_DETENIDA, 3, "· ⛔ La pausa se cortó a los 3 s: el operador pulsó «Detener»."),
                 (m.PAUSA_VENCIDA, 600, "· ⛔ La pausa se cortó a los 600 s: pasaron 10 minutos sin que el operador la "
                                        "resolviera."),
                 (None, 7, "· ▶ La pausa terminó a los 7 s.")]          # el panel Tkinter no dice cómo terminó
        for como, seg, fin in casos:
            with self.subTest(como=como):
                vistas = []
                reg = m.Registro(self.sb / "pausa_panel" / "log.txt", on_log=vistas.append)
                self.addCleanup(reg.cerrar)

                def on_pausa(mensaje, como=como, seg=seg):
                    reloj.t += seg
                    return como
                with mock.patch.object(m, "time", reloj):
                    self.assertEqual(m.pausa_manual(reg, on_pausa, "MSC pide una validación.\nResuélvela."), como)
                self.assertEqual([v.split("] ", 1)[1] for v in vistas],
                                 ["· ⏸ PAUSA: espero al operador. Por qué: MSC pide una validación. Resuélvela.", fin])
        # Si el panel falla, lo anota y la falla sigue su camino.
        vistas = []
        reg = m.Registro(self.sb / "pausa_panel" / "log.txt", on_log=vistas.append)
        self.addCleanup(reg.cerrar)

        def rompe(mensaje):
            raise RuntimeError("el panel se cerró")
        with mock.patch.object(m, "time", reloj), self.assertRaises(RuntimeError):
            m.pausa_manual(reg, rompe, "x")
        self.assertEqual(vistas[-1].split("] ", 1)[1],
                         "· ⛔ La pausa se cortó a los 0 s: falló el panel (RuntimeError: el panel se cerró).")

    def test_argumentos_de_chrome(self):
        base = ["--start-maximized", "--disable-blink-features=AutomationControlled"]
        casos = {None: base + ["--force-device-scale-factor=0.65", "--high-dpi-support=1"],
                 "1": base, "0.999": base, "0.2": base + ["--force-device-scale-factor=0.5", "--high-dpi-support=1"],
                 "0.8": base + ["--force-device-scale-factor=0.8", "--high-dpi-support=1"],
                 "abc": base + ["--force-device-scale-factor=0.65", "--high-dpi-support=1"]}
        for v, esperado in casos.items():
            with self.subTest(zoom=v):
                if v is None:
                    os.environ.pop("AQUASHIELD_ZOOM", None)
                else:
                    os.environ["AQUASHIELD_ZOOM"] = v
                try:
                    self.assertEqual(self.mod._args_chrome(), esperado)
                finally:
                    os.environ.pop("AQUASHIELD_ZOOM", None)

    def test_un_solo_lanzador_del_navegador(self):
        # Todo navegador del programa sale de _lanzar_navegador (encargo 53): el login, las reservas de la consola y las
        # del panel web, y nadie más llama a launch_persistent_context. Así el sandbox de Chrome vale para los tres.
        arbol = ast.parse((soporte.FUENTES / "AQUASHIELD.py").read_text(encoding="utf-8"))
        lanzan, llaman = [], []
        for funcion in ast.walk(arbol):
            if not isinstance(funcion, ast.FunctionDef):
                continue
            for n in ast.walk(funcion):
                if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and \
                        n.func.attr == "launch_persistent_context":
                    lanzan.append(funcion.name)
                if isinstance(n, ast.Call) and isinstance(n.func, ast.Name) and n.func.id == "_lanzar_navegador":
                    llaman.append(funcion.name)
        self.assertEqual(lanzan, ["_lanzar_navegador", "_lanzar_navegador"])          # el canal y el respaldo
        self.assertEqual(sorted(llaman), ["_web_worker", "ejecutar_login", "ejecutar_reservas"])

    def test_navegador_con_el_sandbox_de_chrome(self):
        # Decisión de Marcelo, encargo 53 (CICLO-bloqueo-cma-y-sandbox.md): el navegador se lanza con el sandbox de
        # Chrome activado (chromium_sandbox). Hasta ahí, Playwright le agregaba --no-sandbox porque no se le pedía, y
        # Chrome avisaba de esa bandera en la ventana. Si el canal (Chrome) no abre, el Chromium de Playwright, también
        # con el sandbox, y lo dice en log.txt (hasta el encargo 53, solo el login lo decía; las reservas cambiaban a
        # Chromium sin decirlo). Las demás opciones, las de siempre, y STEALTH_JS como hasta hoy.
        class QueNoAbre(soporte.PlaywrightFalso):
            def launch_persistent_context(self, **kw):
                if "channel" in kw:
                    self.lanzamientos.append(kw)
                    raise RuntimeError("canal falso que no abre")
                return super().launch_persistent_context(**kw)
        perfil = self.sb / "perfiles" / "op_navegador"
        opciones = dict(user_data_dir=str(perfil), headless=False, no_viewport=True, args=self.mod._args_chrome(),
                        ignore_default_args=["--enable-automation"], chromium_sandbox=True)
        for pw, canales in ((soporte.PlaywrightFalso(), ["chrome"]), (QueNoAbre(), ["chrome", None])):
            with self.subTest(canales=canales):
                ruta = self.sb / "logs" / f"navegador_{len(canales)}" / "log.txt"
                reg = self.mod.Registro(ruta)
                try:
                    ctx = self.mod._lanzar_navegador(pw, perfil, False, "chrome", reg)
                finally:
                    reg.cerrar()
                self.assertEqual([kw.get("channel") for kw in pw.lanzamientos], canales)
                self.assertEqual([{k: v for k, v in kw.items() if k != "channel"} for kw in pw.lanzamientos],
                                 [opciones] * len(canales))
                self.assertEqual(ctx.scripts, [self.mod.STEALTH_JS])
                log = ruta.read_text(encoding="utf-8")
                cambio = "no pude usar 'chrome' (canal falso que no abre); uso chromium de playwright"
                self.assertEqual(cambio in log, len(canales) == 2)

    def test_volcado_html_cae_en_la_carpeta_del_programa(self):
        # Modo descubrir: los volcados de página van a la raíz del programa (BASE), junto a
        # la planilla y la config; por eso el .gitignore excluye *_dump.html.
        reg = self.mod.Registro(self.sb / "volcado" / "log.txt")
        pagina = soporte.PaginaFalsa()
        pagina.content = lambda: "<html>volcado sintetico</html>"
        self.mod._guardar_html(pagina, "msc_dump.html", reg)

        class Rompe(soporte.PaginaFalsa):
            def content(self):
                raise RuntimeError("pagina cerrada")
        self.mod._guardar_html(Rompe(), "otro_dump.html", reg)
        reg.cerrar()
        self.assertEqual((self.sb / "msc_dump.html").read_text(encoding="utf-8"), "<html>volcado sintetico</html>")
        self.assertFalse((self.sb / "otro_dump.html").exists())
        self.assertIn("no pude guardar otro_dump.html: pagina cerrada",
                      (self.sb / "volcado" / "log.txt").read_text(encoding="utf-8"))
        (self.sb / "msc_dump.html").unlink()

    def test_config_ausente(self):
        ruta = self.sb / "config.json"
        original = ruta.read_bytes()
        ruta.unlink()
        try:
            with self.assertRaises(FileNotFoundError) as e:
                self.mod.cargar_config()
        finally:
            ruta.write_bytes(original)
        self.assertIn("copia config.example.json", str(e.exception))


class TestLeeme(unittest.TestCase):
    """LEEME.md describe el programa real. Hasta 42247d4 hablaba de AQUASHIELD.exe, crear_exe.bat e «Iniciar
    AQUASHIELD (respaldo).bat», que no existen (CICLO-cola-nueve-items.md)."""

    def test_nombra_solo_archivos_del_programa_que_existen(self):
        # De la copia de las fuentes si la trae (el censo deja ahí el LEEME mutado); si no, del repo.
        ruta = soporte.FUENTES / "LEEME.md" if (soporte.FUENTES / "LEEME.md").exists() else soporte.RAIZ / "LEEME.md"
        texto = ruta.read_text(encoding="utf-8")
        nombrados = sorted(set(re.findall(r"`([^`]+\.(?:bat|exe|py))`", texto)))
        self.assertEqual(nombrados, ["AQUASHIELD.py", "AQUASHIELD_EMISION.py", "Iniciar AQUASHIELD.bat",
                                     "Iniciar AQUASHIELD_EMISION.bat"])
        self.assertEqual([n for n in nombrados if not (soporte.FUENTES / n).exists()], [])
        self.assertNotIn(".exe", texto)


class TestLanzadores(unittest.TestCase):
    """Los dos lanzadores instalan holidays si falta, antes de abrir el panel: los feriados de Chile de la fecha de
    retiro de MAERSK (decisión de Marcelo, CICLO-maersk-cuatro-puntos.md). Sin red no frenan el arranque: con
    --timeout 5 y --retries 0, pip se rindió en 6 s ante un índice que no responde, y sin esos límites en 69 s, con 5
    reintentos (medido el 2026-09-29). Si igual no queda, el programa avisa y sigue de lunes a viernes."""
    LINEA = ('python -c "import holidays" 2>nul || python -m pip install --user --disable-pip-version-check '
             '--timeout 5 --retries 0 holidays')

    def test_instalan_holidays_sin_frenar_el_arranque(self):
        for nombre, programa in (("Iniciar AQUASHIELD.bat", "AQUASHIELD.py"),
                                 ("Iniciar AQUASHIELD_EMISION.bat", "AQUASHIELD_EMISION.py")):
            with self.subTest(lanzador=nombre):
                # De soporte.FUENTES: el censo deja ahí el .bat mutado.
                lineas = (soporte.FUENTES / nombre).read_text(encoding="utf-8").split("\n")
                instala = [i for i, l in enumerate(lineas) if "holidays" in l and not l.upper().startswith("REM")]
                abre = [i for i, l in enumerate(lineas) if l == f'start "" pythonw "{programa}"']
                self.assertEqual([lineas[i] for i in instala], [self.LINEA])
                self.assertEqual(len(abre), 1)
                self.assertLess(instala[0], abre[0])
                # Entre la instalación y el panel, nada que lo detenga.
                entre = [l for l in lineas[instala[0] + 1:abre[0]] if l.strip() and not l.upper().startswith("REM")]
                self.assertEqual([l for l in entre if re.search(r"\b(exit|pause|goto|errorlevel)\b", l, re.I)], [])


if __name__ == "__main__":
    unittest.main()
