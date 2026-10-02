# -*- coding: utf-8 -*-
"""Casos sintéticos: la ESTRUCTURA de los datos reales con valores inventados.

La planilla copia la forma de la operativa (7 hojas, títulos en las filas 1-2
combinadas, fila 3 vacía, encabezados en la fila 4 con los mismos textos; en
COSCO la columna H se llama «Estado (Robot)» y no hay columna I). Los valores
(puertos, naves, viajes, contratos, usuarios, claves) son inventados. Se genera
al correr: no se versiona ningún .xlsx.
"""
import openpyxl

NAVIERAS = ("one", "msc", "cma", "cosco", "hyundai", "maersk")

ENCABEZADOS = ["Naviera", "Pto. Embarque", "Puerto Descarga", "Destino Final", "NAVE", "Viaje",
               "Contrato / Cotización", "Estado", "N° de reserva emitida"]
ENCABEZADOS_COSCO = ENCABEZADOS[:7] + ["Estado (Robot)"]

TITULO_1 = "PLANILLA SINTETICA DE PRUEBA - reservas de ejemplo"
TITULO_2 = "Valores inventados; misma estructura que la planilla operativa"
TITULO_2_TRAMPA = "Valores inventados; puerto de origen y destino segun la hoja"

# Filas de datos (desde la fila 5). None = celda vacía.
FILAS = {
    "ONE": [
        ["ONE", "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA UNO", "V001E", "CT-PRUEBA-ONE", None, None],
        ["ONE", "PUERTO ALFA", "PUERTO DELTA", None, "NAVE PRUEBA DOS", "V002E", None, None, None],
        ["ONE", "PUERTO ALFA", "PUERTO EPSILON", "CIUDAD ZETA", None, "ONE PRUEBA 003W", None, None, None],
        [None] * 9,
        [None, "elegir siempre la primera nave disponible", None, None, None, None, None, None, None],
        ["ONE", "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE POR COMPLETAR", "V004E", None, None, None],
    ],
    "MSC": [
        ["MSC", "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA TRES", "V010N", "CT-PRUEBA-MSC", None, None],
        ["MSC", "PUERTO ALFA", "PUERTO DELTA", "CIUDAD ETA", None, "MSC PRUEBA NX1", None, None, None],
        ["MSC", "PUERTO ALFA", None, "CIUDAD THETA", "NAVE PRUEBA CUATRO", "V011N", None, None, None],
        ["Naviera", "Pto. Embarque", "Puerto Descarga", "Destino Final", "NAVE", "Viaje", None, None, None],
    ],
    "CMA-CGM": [
        ["CMA CGM", "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA CINCO", "V020S", "CT-PRUEBA-CMA", None, None],
        ["CMA CGM", "PUERTO ALFA", "PUERTO IOTA", "CIUDAD KAPPA", "NAVE PRUEBA SEIS", "V021S", None, None, None],
        [None, None, None, None, "NAVE PRUEBA, VER NOTA", None, None, None, None],
    ],
    "COSCO": [
        ["COSCO", "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA SIETE", "V030W", "CT-PRUEBA-COSCO", None, None],
        ["COSCO", "PUERTO ALFA", "PUERTO LAMBDA", "CIUDAD MU", None, None, None, None, None],
    ],
    "HYUNDAI": [
        ["HYUNDAI", "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA OCHO", "V040E", "CT-PRUEBA-HMM", None, None],
        ["HMM", "PUERTO ALFA", "PUERTO NU", "CIUDAD XI", "NAVE PRUEBA NUEVE", "V041E", None, None, None],
    ],
    "MAERSK": [
        ["MAERSK", "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA DIEZ", "V050N", "CT-PRUEBA-MAERSK", None, None],
        ["Maersk Line", "PUERTO ALFA", "PUERTO OMICRON", "CIUDAD PI", "NAVE PRUEBA ONCE", "V051N", None, None, None],
    ],
    "CONSOLIDADO": [
        ["ONE", "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA UNO", "V001E", "CT-PRUEBA-ONE", None, None],
        ["MSC", "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA TRES", "V010N", "CT-PRUEBA-MSC", None, None],
        ["CMA CGM", "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA CINCO", "V020S", None, None, None],
        ["COSCO", "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA SIETE", "V030W", None, None, None],
        ["HMM", "PUERTO ALFA", "PUERTO NU", "CIUDAD XI", "NAVE PRUEBA NUEVE", "V041E", None, None, None],
        ["Maersk Line", "PUERTO ALFA", "PUERTO OMICRON", "CIUDAD PI", "NAVE PRUEBA ONCE", "V051N", None, None, None],
        ["OTRA LINEA", "PUERTO ALFA", "PUERTO RHO", "CIUDAD SIGMA", "NAVE PRUEBA DOCE", "V060X", None, None, None],
    ],
}
ORDEN_HOJAS = ["CONSOLIDADO", "ONE", "MSC", "CMA-CGM", "COSCO", "HYUNDAI", "MAERSK"]


def config(emitir=False):
    """config.json sintético: dos operadores, credenciales inventadas, llave de config apagada."""
    def cred(nav, llena=True):
        return {"usuario": f"usuario.{nav}@ejemplo.invalid" if llena else "",
                "clave": f"clave-falsa-{nav}" if llena else "",
                "contrato": f"CT-PRUEBA-{nav.upper()}" if llena else ""}
    return {
        "_comentario": "config sintetica de la red de verificacion",
        "opciones": {"headless": False, "navegador": "chrome", "espera_seg": 30, "emitir_reservas": emitir,
                     "peso_reefer_kg": 22500, "temperatura_reefer_c": -20},
        "usuarios": {
            "op_prueba": {"descripcion": "Operador de prueba", **{n: cred(n) for n in NAVIERAS}},
            "op_vacio": {"descripcion": "Operador sin credenciales", **{n: cred(n, False) for n in NAVIERAS}},
        },
    }


def crear_planilla(ruta, trampa_titulo=False, hojas=None):
    """Escribe la planilla sintética en 'ruta'. trampa_titulo: el título de la
    hoja ONE nombra «origen», una de las palabras con que el lector de consola
    reconoce la fila de encabezados."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre in hojas or ORDEN_HOJAS:
        ws = wb.create_sheet(nombre)
        ws.cell(1, 1, TITULO_1)
        ws.cell(2, 1, TITULO_2_TRAMPA if (trampa_titulo and nombre == "ONE") else TITULO_2)
        ws.merge_cells("A1:I1")
        ws.merge_cells("A2:I2")
        for c, h in enumerate(ENCABEZADOS_COSCO if nombre == "COSCO" else ENCABEZADOS, start=1):
            ws.cell(4, c, h)
        for i, fila in enumerate(FILAS[nombre], start=5):
            for c, v in enumerate(fila, start=1):
                if v is not None:
                    ws.cell(i, c, v)
    wb.save(ruta)
    return ruta
