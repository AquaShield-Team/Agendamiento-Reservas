# -*- coding: utf-8 -*-
"""Los dos lectores de la planilla, sus diferencias y cómo escribe cada uno.

- Consola: leer_reservas(naviera) lee la hoja de la naviera desde
  «Reservas AQUASHIELD.xlsx» junto al programa; escribir_estado() escribe
  el resultado SOBRE ese mismo archivo.
- Panel web: leer_filas(hoja) lee cualquier hoja ya abierta; la descarga
  (_planilla_con_resultados) devuelve una COPIA y no toca el archivo.
Todo sobre la planilla sintética, en la sandbox. Las diferencias quedan
fotografiadas tal como están: no se corrigen.
"""
import hashlib
import shutil
import unittest
from io import BytesIO

import openpyxl

import sinteticos
import soporte

# La foto del lector de la consola. En ONE (fila 7) y MSC (fila 6) la celda de la nave viene vacía y la columna F
# («Viaje») trae una palabra de su lista: desde CICLO-cola-seis-items.md la nave queda vacía, y lo demás igual (en ONE, el
# destino final sigue siendo el puerto de descarga, «PUERTO EPSILON», no la columna Destino Final). Hasta f68ce6c la
# nave era el viaje: «ONE PRUEBA 003W» y «MSC PRUEBA NX1».
CONSOLA = {
    "one": [(5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA UNO", "V001E", "CT-PRUEBA-ONE"),
            (6, "PUERTO ALFA", "PUERTO DELTA", "PUERTO DELTA", "NAVE PRUEBA DOS", "V002E", ""),
            (7, "PUERTO ALFA", "PUERTO EPSILON", "PUERTO EPSILON", "", "ONE PRUEBA 003W", ""),
            (10, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE POR COMPLETAR", "V004E", "")],
    "msc": [(5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA TRES", "V010N", "CT-PRUEBA-MSC"),
            (6, "PUERTO ALFA", "PUERTO DELTA", "CIUDAD ETA", "", "MSC PRUEBA NX1", ""),
            (7, "PUERTO ALFA", "CIUDAD THETA", "CIUDAD THETA", "NAVE PRUEBA CUATRO", "V011N", ""),
            (8, "Pto. Embarque", "Puerto Descarga", "Destino Final", "NAVE", "Viaje", "")],
    "cma": [(5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA CINCO", "V020S", "CT-PRUEBA-CMA"),
            (6, "PUERTO ALFA", "PUERTO IOTA", "CIUDAD KAPPA", "NAVE PRUEBA SEIS", "V021S", "")],
    "cosco": [(5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA SIETE", "V030W", "CT-PRUEBA-COSCO"),
              (6, "PUERTO ALFA", "PUERTO LAMBDA", "CIUDAD MU", "", "", "")],
    "hyundai": [(5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA OCHO", "V040E", "CT-PRUEBA-HMM"),
                (6, "PUERTO ALFA", "PUERTO NU", "CIUDAD XI", "NAVE PRUEBA NUEVE", "V041E", "")],
    "maersk": [(5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA DIEZ", "V050N", "CT-PRUEBA-MAERSK"),
               (6, "PUERTO ALFA", "PUERTO OMICRON", "CIUDAD PI", "NAVE PRUEBA ONCE", "V051N", "")],
}
WEB = {
    "CONSOLIDADO": [("ONE", 5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA UNO", "V001E", "CT-PRUEBA-ONE"),
                    ("MSC", 6, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA TRES", "V010N", "CT-PRUEBA-MSC"),
                    ("CMA CGM", 7, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA CINCO", "V020S", ""),
                    ("COSCO", 8, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA SIETE", "V030W", ""),
                    ("HMM", 9, "PUERTO ALFA", "PUERTO NU", "CIUDAD XI", "NAVE PRUEBA NUEVE", "V041E", ""),
                    ("Maersk Line", 10, "PUERTO ALFA", "PUERTO OMICRON", "CIUDAD PI", "NAVE PRUEBA ONCE", "V051N", ""),
                    ("OTRA LINEA", 11, "PUERTO ALFA", "PUERTO RHO", "CIUDAD SIGMA", "NAVE PRUEBA DOCE", "V060X", "")],
    "ONE": [("ONE", 5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA UNO", "V001E", "CT-PRUEBA-ONE"),
            ("ONE", 6, "PUERTO ALFA", "PUERTO DELTA", "PUERTO DELTA", "NAVE PRUEBA DOS", "V002E", ""),
            ("ONE", 7, "PUERTO ALFA", "PUERTO EPSILON", "CIUDAD ZETA", "", "ONE PRUEBA 003W", ""),
            ("", 9, "elegir siempre la primera nave disponible", "", "", "", "", ""),
            ("ONE", 10, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "", "V004E", "")],
    "MSC": [("MSC", 5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA TRES", "V010N", "CT-PRUEBA-MSC"),
            ("MSC", 6, "PUERTO ALFA", "PUERTO DELTA", "CIUDAD ETA", "", "MSC PRUEBA NX1", ""),
            ("MSC", 7, "PUERTO ALFA", "", "CIUDAD THETA", "NAVE PRUEBA CUATRO", "V011N", "")],
    "CMA-CGM": [("CMA CGM", 5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA CINCO", "V020S", "CT-PRUEBA-CMA"),
                ("CMA CGM", 6, "PUERTO ALFA", "PUERTO IOTA", "CIUDAD KAPPA", "NAVE PRUEBA SEIS", "V021S", ""),
                ("", 7, "", "", "", "NAVE PRUEBA, VER NOTA", "", "")],
    "COSCO": [("COSCO", 5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA SIETE", "V030W", "CT-PRUEBA-COSCO"),
              ("COSCO", 6, "PUERTO ALFA", "PUERTO LAMBDA", "CIUDAD MU", "", "", "")],
    "HYUNDAI": [("HYUNDAI", 5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA OCHO", "V040E", "CT-PRUEBA-HMM"),
                ("HMM", 6, "PUERTO ALFA", "PUERTO NU", "CIUDAD XI", "NAVE PRUEBA NUEVE", "V041E", "")],
    "MAERSK": [("MAERSK", 5, "PUERTO ALFA", "PUERTO BETA", "CIUDAD GAMMA", "NAVE PRUEBA DIEZ", "V050N", "CT-PRUEBA-MAERSK"),
               ("Maersk Line", 6, "PUERTO ALFA", "PUERTO OMICRON", "CIUDAD PI", "NAVE PRUEBA ONCE", "V051N", "")],
}
CAMPOS = ("pol", "destino_orig", "destino_final", "nave", "viaje", "cotizacion")
CLAVES_CONSOLA = {"fila", "pol", "reserva", "destino_orig", "destino_final", "dia_retiro", "deposito",
                  "dia_carga", "nave", "viaje", "contenedor", "cant", "cotizacion"}
COLUMNAS_WEB = {"naviera": 1, "pol": 2, "destino_orig": 3, "destino_final": 4, "nave": 5, "viaje": 6,
                "cotizacion": 7, "estado": 8, "reserva_emitida": 9}


def sha(ruta):
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


def foto_consola(filas):
    return [(f["fila"],) + tuple(f[c] for c in CAMPOS) for f in filas]


def foto_web(filas):
    return [(f["naviera"], f["fila"]) + tuple(f[c] for c in CAMPOS) for f in filas]


class ConPlanilla(soporte.CasoAQ):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.planilla = sinteticos.crear_planilla(cls.sb / "Reservas AQUASHIELD.xlsx")

    def setUp(self):
        super().setUp()
        sinteticos.crear_planilla(self.planilla)      # cada prueba parte de la planilla limpia


class TestLectorConsola(ConPlanilla):
    def _nav(self, nav):
        filas = self.mod.leer_reservas(nav)
        self.assertEqual(foto_consola(filas), CONSOLA[nav])
        for f in filas:
            self.assertEqual(set(f), CLAVES_CONSOLA)
            self.assertEqual((f["reserva"], f["contenedor"], f["cant"], f["dia_carga"]), ("", "40 HC REEFER", 1, ""))

    def test_consola_one(self):
        self._nav("one")

    def test_consola_msc(self):
        self._nav("msc")

    def test_consola_cma(self):
        self._nav("cma")

    def test_consola_cosco(self):
        self._nav("cosco")

    def test_consola_hyundai(self):
        self._nav("hyundai")

    def test_consola_maersk(self):
        self._nav("maersk")

    def test_consola_trampa_de_titulo(self):
        # Si el título nombra «origen», la consola toma esa fila como encabezado: lee
        # la columna Naviera como puerto de embarque y deja lo demás vacío. La regla de
        # ONE mira la columna 6 fija, pero ya no le da la nave: hasta f68ce6c recuperaba
        # «ONE PRUEBA 003W» en la fila 7 (CICLO-cola-seis-items.md).
        sinteticos.crear_planilla(self.planilla, trampa_titulo=True)
        filas = self.mod.leer_reservas("one")
        self.assertEqual([(f["fila"], f["pol"], f["nave"], f["destino_final"]) for f in filas],
                         [(4, "Naviera", "", ""), (5, "ONE", "", ""), (6, "ONE", "", ""),
                          (7, "ONE", "", ""), (10, "ONE", "", "")])

    def test_la_f_sigue_decidiendo_las_filas(self):
        # En ONE y MSC, la columna F ya no da la nave, pero sigue decidiendo qué filas lee la consola, como hasta f68ce6c
        # (CICLO-cola-seis-items.md): una fila con solo la naviera y la F entra a la lista, sin nave; y una nota escrita en
        # la F («no hacerla aun», «elegir siempre») se descarta.
        wb = openpyxl.load_workbook(self.planilla)
        for hoja, f in (("ONE", "ONE PRUEBA 004W"), ("MSC", "MSC PRUEBA NX2")):
            wb[hoja]["A20"], wb[hoja]["F20"] = hoja, f
            wb[hoja]["B21"], wb[hoja]["F21"] = "PUERTO ALFA", f"{hoja}: no hacerla aun, elegir siempre la primera"
        wb.save(self.planilla)
        for nav in ("one", "msc"):
            with self.subTest(naviera=nav):
                filas = {f["fila"]: f for f in self.mod.leer_reservas(nav)}
                self.assertEqual((20 in filas, filas.get(20, {}).get("nave"), 21 in filas), (True, "", False))

    def test_consola_sin_archivo(self):
        self.planilla.unlink()
        with self.assertRaises(FileNotFoundError) as e:
            self.mod.leer_reservas("one")
        self.assertIn("Falta el archivo Excel de reservas", str(e.exception))

    def test_consola_hoja_faltante(self):
        sinteticos.crear_planilla(self.planilla, hojas=["CONSOLIDADO", "ONE"])
        with self.assertRaises(ValueError) as e:
            self.mod.leer_reservas("maersk")
        self.assertIn("no tiene la hoja", str(e.exception))

    def test_encontrar_archivo_prioridad(self):
        nombres = ("Reservas AQUASHIELD.xlsx", "Planilla_Estandar_AQUASHIELD.xlsx", "Agendamientos.xlsx")
        self.planilla.unlink()
        self.assertEqual(self.mod.encontrar_archivo_reservas().name, nombres[0])   # por defecto, aunque falte
        for n in reversed(nombres):
            (self.sb / n).write_bytes(b"")
            self.assertEqual(self.mod.encontrar_archivo_reservas().name, n)
        for n in nombres[1:]:
            (self.sb / n).unlink()


class TestLectorWeb(ConPlanilla):
    def _hoja(self, hoja, columnas=None):
        ws = openpyxl.load_workbook(self.planilla, data_only=True)[hoja]
        fila_enc, cols, filas = self.mod.leer_filas(ws)
        self.assertEqual(fila_enc, 4)
        self.assertEqual(cols, columnas or COLUMNAS_WEB)
        self.assertEqual(foto_web(filas), WEB[hoja])
        for f in filas:
            self.assertEqual(set(f), CLAVES_CONSOLA | {"naviera", "temp"})

    def test_web_consolidado(self):
        self._hoja("CONSOLIDADO")

    def test_web_one(self):
        self._hoja("ONE")

    def test_web_msc(self):
        self._hoja("MSC")

    def test_web_cma(self):
        self._hoja("CMA-CGM")

    def test_web_cosco(self):
        cols = {k: v for k, v in COLUMNAS_WEB.items() if k != "reserva_emitida"}
        self._hoja("COSCO", cols)

    def test_web_hyundai(self):
        self._hoja("HYUNDAI")

    def test_web_maersk(self):
        self._hoja("MAERSK")

    def test_web_trampa_de_titulo(self):
        # El panel web elige la fila con MÁS encabezados reconocidos: el título no lo engaña.
        sinteticos.crear_planilla(self.planilla, trampa_titulo=True)
        self._hoja("ONE")


class TestDiferencias(ConPlanilla):
    """Las diferencias entre los dos lectores sobre la MISMA hoja, fotografiadas."""
    ESPERADO = {
        # Hasta f68ce6c la consola también daba otra nave en la 7 de ONE y en la 6 de MSC: el viaje, de la columna F
        # (CICLO-cola-seis-items.md).
        "one": ([], [9], {7: {"destino_final": ("PUERTO EPSILON", "CIUDAD ZETA")},
                          10: {"nave": ("NAVE POR COMPLETAR", "")}}),
        "msc": ([8], [], {7: {"destino_orig": ("CIUDAD THETA", "")}}),
        "cma": ([], [7], {}),
        "cosco": ([], [], {}),
        "hyundai": ([], [], {}),
        "maersk": ([], [], {}),
    }
    HOJA = {"one": "ONE", "msc": "MSC", "cma": "CMA-CGM", "cosco": "COSCO", "hyundai": "HYUNDAI", "maersk": "MAERSK"}

    def test_diferencias_entre_lectores(self):
        wb = openpyxl.load_workbook(self.planilla, data_only=True)
        for nav, hoja in self.HOJA.items():
            con = {f["fila"]: f for f in self.mod.leer_reservas(nav)}
            web = {f["fila"]: f for f in self.mod.leer_filas(wb[hoja])[2]}
            solo_con = sorted(set(con) - set(web))
            solo_web = sorted(set(web) - set(con))
            distintos = {}
            for fila in sorted(set(con) & set(web)):
                d = {c: (con[fila][c], web[fila][c]) for c in CAMPOS if con[fila][c] != web[fila][c]}
                if d:
                    distintos[fila] = d
            with self.subTest(naviera=nav):
                self.assertEqual((solo_con, solo_web, distintos), self.ESPERADO[nav])
                self.assertEqual(set(web[min(web)]) - set(con[min(con)]), {"naviera", "temp"})


class TestEscritura(ConPlanilla):
    def _valores(self, ruta):
        wb = openpyxl.load_workbook(ruta)
        return {(ws.title, c.coordinate): c.value for ws in wb.worksheets for fila in ws.iter_rows() for c in fila}

    def test_consola_escribe_sobre_la_planilla(self):
        antes_sha, antes = sha(self.planilla), self._valores(self.planilla)
        self.assertIsNone(self.mod.escribir_estado("one", 6, "ESTADO DE PRUEBA"))
        self.assertNotEqual(sha(self.planilla), antes_sha, "escribir_estado no modificó el archivo original")
        despues = self._valores(self.planilla)
        cambios = {k: (antes.get(k), despues.get(k)) for k in set(antes) | set(despues) if antes.get(k) != despues.get(k)}
        self.assertEqual(cambios, {("ONE", "H6"): (None, "ESTADO DE PRUEBA")})

    def test_consola_cosco_usa_estado_robot(self):
        self.mod.escribir_estado("cosco", 5, "X")
        ws = openpyxl.load_workbook(self.planilla)["COSCO"]
        self.assertEqual((ws["H4"].value, ws["H5"].value), ("Estado (Robot)", "X"))

    def test_consola_sin_columna_estado_crea_una(self):
        # Sin columna «Estado», crea «Estado (Robot)» en la siguiente a la última con
        # encabezado, nunca antes de la M. Hasta d86fbca caía en la P (16): el bucle que la
        # buscaba usaba ws.cell(), que CREA celdas, y cada fila revisada corría max_column.
        wb = openpyxl.load_workbook(self.planilla)
        wb["ONE"]["H4"] = "Observaciones"
        wb.save(self.planilla)
        self.mod.escribir_estado("one", 5, "X")
        ws = openpyxl.load_workbook(self.planilla)["ONE"]
        self.assertEqual((ws["M4"].value, ws["M5"].value, ws["P4"].value, ws["H5"].value),
                         ("Estado (Robot)", "X", None, None))

    def test_consola_columna_nueva_no_depende_de_max_column(self):
        # Una nota suelta en N7, sin encabezado arriba, corre max_column a 14: la columna
        # nueva sigue en la M, porque se ubica por los encabezados (el último está en la I).
        wb = openpyxl.load_workbook(self.planilla)
        wb["ONE"]["H4"] = "Observaciones"
        wb["ONE"]["N7"] = "nota suelta de prueba"
        wb.save(self.planilla)
        self.mod.escribir_estado("one", 5, "X")
        ws = openpyxl.load_workbook(self.planilla)["ONE"]
        self.assertEqual((ws["M4"].value, ws["M5"].value, ws["O4"].value, ws["O5"].value),
                         ("Estado (Robot)", "X", None, None))

    def test_consola_encabezado_en_otra_fila(self):
        # Con los encabezados en la fila 5 y sin «Estado», «Estado (Robot)» nace en la fila
        # de encabezados, no en la 4 fija de antes.
        wb = openpyxl.load_workbook(self.planilla)
        ws = wb["ONE"]
        for rango in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(rango))
        ws.insert_rows(1)
        ws["H5"] = "Observaciones"
        wb.save(self.planilla)
        self.mod.escribir_estado("one", 6, "X")
        ws = openpyxl.load_workbook(self.planilla)["ONE"]
        self.assertEqual((ws["M5"].value, ws["M6"].value, ws["M4"].value), ("Estado (Robot)", "X", None))

    def test_consola_error_de_escritura_se_levanta(self):
        # Sin planilla, escribir_estado levanta ErrorEscritura con el motivo en español
        # (hasta fba3bc3 devolvía None sin avisar).
        self.planilla.unlink()
        with self.assertRaises(self.mod.ErrorEscritura) as e:
            self.mod.escribir_estado("one", 5, "X")
        self.assertEqual(str(e.exception), "no encontré «Reservas AQUASHIELD.xlsx» junto al programa")
        self.assertFalse(self.planilla.exists())

    def test_consola_errores_de_escritura_en_espanol(self):
        casos = []
        sinteticos.crear_planilla(self.planilla, hojas=["CONSOLIDADO", "ONE"])
        casos.append(("maersk", 5, "«Reservas AQUASHIELD.xlsx» no tiene la hoja de maersk (busqué: MAERSK)"))
        wb = openpyxl.load_workbook(self.planilla)
        ws = wb["ONE"]                           # hoja vacía: ningún encabezado que reconocer
        for rango in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(rango))
        ws.delete_rows(1, ws.max_row)
        wb.save(self.planilla)
        casos.append(("one", 5, "no encontré la fila de encabezados en la hoja «ONE»"))
        for nav, fila, motivo in casos:
            with self.subTest(naviera=nav):
                with self.assertRaises(self.mod.ErrorEscritura) as e:
                    self.mod.escribir_estado(nav, fila, "X")
                self.assertEqual(str(e.exception), motivo)
        # Cualquier otro error también se levanta: la fila 1 es un título combinado (A1:I1).
        sinteticos.crear_planilla(self.planilla)
        with self.assertRaises(self.mod.ErrorEscritura) as e:
            self.mod.escribir_estado("one", 1, "X")
        self.assertTrue(str(e.exception).startswith(
            "error inesperado al escribir en «Reservas AQUASHIELD.xlsx» (AttributeError: "), str(e.exception))

    def _cambios_de_la_descarga(self, hoja, resultados):
        antes = self._valores(self.planilla)
        datos = self.mod._planilla_con_resultados(str(self.planilla), hoja, resultados)
        despues = self._valores(BytesIO(datos))
        return {k: (antes.get(k), despues.get(k)) for k in set(antes) | set(despues) if antes.get(k) != despues.get(k)}

    def test_web_descarga_es_copia_y_solo_escribe_la_hoja_corrida(self):
        # Solo la hoja que se corrió. Hasta 4f5f355 los mismos números de fila se escribían
        # también en CONSOLIDADO, donde la fila 7 es otra reserva (en la sintética, de CMA CGM).
        antes = sha(self.planilla)
        resultados = {"5": {"estado": "EMITIDA", "detalle": "x", "booking": "BKGFALSO01"},
                      "7": {"estado": "OK-EJEMPLO", "detalle": "armada sin emitir", "booking": ""}}
        cambios = self._cambios_de_la_descarga("ONE", resultados)
        self.assertEqual(sha(self.planilla), antes, "la descarga modificó el archivo")
        # La OK-EJEMPLO no se emitió: su «N° de reserva emitida» dice que quedó lista para emitir, como el panel. Hasta el
        # encargo 34 quedaba vacía, y el panel decía «✓ Confirmada» (CICLO-one-y-maersk-eligen-la-nave.md).
        self.assertEqual(cambios, {("ONE", "H5"): (None, "EMITIDA"), ("ONE", "I5"): (None, "BKGFALSO01"),
                                   ("ONE", "H7"): (None, "OK-EJEMPLO · armada sin emitir"),
                                   ("ONE", "I7"): (None, "lista para emitir")})

    def test_web_descarga_de_consolidado_solo_escribe_consolidado(self):
        # Al correr CONSOLIDADO, los números de fila son de CONSOLIDADO: no se reparten a las hojas.
        cambios = self._cambios_de_la_descarga(
            "CONSOLIDADO", {"6": {"estado": "EMITIDA", "detalle": "x", "booking": "BKGFALSO04"}})
        self.assertEqual(cambios, {("CONSOLIDADO", "H6"): (None, "EMITIDA"),
                                   ("CONSOLIDADO", "I6"): (None, "BKGFALSO04")})

    def test_web_descarga_avisa_lo_que_no_escribe(self):
        # Lo que no se pudo escribir sale en el registro del panel y en el log.txt de la corrida,
        # con el texto para anotarlo a mano, igual que la consola (hasta 0242321 se callaba).
        carpeta = self.sb / "logs" / "web_one_prueba_descarga"
        emitida = {"estado": "EMITIDA", "detalle": "x", "booking": "BKGFALSO06"}
        armada = {"estado": "OK-EJEMPLO", "detalle": "armada"}
        vacia = sinteticos.crear_planilla(self.sb / "sin_encabezados.xlsx")
        wb = openpyxl.load_workbook(vacia)
        for rango in list(wb["ONE"].merged_cells.ranges):
            wb["ONE"].unmerge_cells(str(rango))
        wb["ONE"].delete_rows(1, wb["ONE"].max_row)
        wb.save(vacia)
        casos = [  # (planilla, hoja, resultados, motivo de la fila avisada, fila avisada)
            (self.planilla, "ONE", {"1": emitida, "5": armada}, "error inesperado (AttributeError: ", "1"),
            (self.planilla, "NO-EXISTE", {"5": emitida}, "la planilla no tiene la hoja «NO-EXISTE»", "5"),
            (vacia, "ONE", {"5": emitida}, "no encontré la fila de encabezados en la hoja «ONE»", "5"),
        ]
        self.mod._WEB["carpeta"] = str(carpeta)
        descargas = []
        try:
            for ruta, hoja, resultados, motivo, fila in casos:
                with self.subTest(hoja=hoja, motivo=motivo[:20]):
                    desde = len(self.mod._WEB["log"])
                    descargas.append(self.mod._planilla_con_resultados(str(ruta), hoja, resultados))
                    nuevas = self.mod._WEB["log"][desde:]
                    log = (carpeta / "log.txt").read_text(encoding="utf-8").splitlines()[-2:]
                    aviso = f"· ⚠ No pude escribir en la planilla descargada el resultado de la fila {fila}: {motivo}"
                    resumen = (f"· ⚠ 1 resultado(s) no quedaron en la planilla descargada (filas {fila}); "
                               f"cada uno está arriba en este log.")
                    self.assertEqual([aviso in log[0], resumen in log[1]], [True, True], log)
                    self.assertEqual([aviso in nuevas[0], resumen in nuevas[1], len(nuevas)], [True, True, 2], nuevas)
                    self.assertIn("Anótalo a mano: " + ("EMITIDA · BKGFALSO06 · x" if resultados[fila] is emitida
                                                        else "OK-EJEMPLO · armada"), log[0])
            # En el primer caso, la fila 5 de ONE, que sí se podía, quedó escrita junto al aviso de la fila 1.
            self.assertEqual(openpyxl.load_workbook(BytesIO(descargas[0]))["ONE"]["H5"].value, "OK-EJEMPLO · armada")
            # Sin corrida (sin carpeta), el aviso igual llega al panel.
            self.mod._WEB["carpeta"] = ""
            desde = len(self.mod._WEB["log"])
            self.mod._planilla_con_resultados(str(self.planilla), "NO-EXISTE", {"5": armada})
            self.assertEqual(len(self.mod._WEB["log"]) - desde, 2)
            self.assertIn("la planilla no tiene la hoja «NO-EXISTE». Anótalo a mano: OK-EJEMPLO · armada",
                          self.mod._WEB["log"][desde])
        finally:
            self.mod._WEB["carpeta"] = ""

    def test_web_descarga_de_hoja_que_no_esta_no_escribe(self):
        # Si la hoja corrida no está en el libro, no se escribe nada (hasta 4f5f355, en la primera hoja).
        cambios = self._cambios_de_la_descarga(
            "NO-EXISTE", {"5": {"estado": "EMITIDA", "detalle": "x", "booking": "BKGFALSO05"}})
        self.assertEqual(cambios, {})

    def test_web_descarga_crea_columna_de_numero(self):
        datos = self.mod._planilla_con_resultados(str(self.planilla), "COSCO",
                                                  {"5": {"estado": "EMITIDA", "booking": "BKGFALSO02"}})
        ws = openpyxl.load_workbook(BytesIO(datos))["COSCO"]
        self.assertEqual((ws["H4"].value, ws["I4"].value), ("Estado (Robot)", "N° de reserva emitida"))
        self.assertEqual((ws["H5"].value, ws["I5"].value), ("EMITIDA", "BKGFALSO02"))


class TestEstadosNuevos(ConPlanilla):
    """Los estados de después del envío, escritos en la columna de estado como los escriben la
    consola y la descarga del panel, no cambian lo que leen los dos lectores: ninguno lee el
    valor de esa columna (medido para CICLO-estado-tras-envio.md)."""
    ESTADOS = [("ENVIADA – REVISAR EN PORTAL", "ONE: se pulsó «Submit», pero la pantalla no mostró la "
                "confirmación con el número de reserva."),
               ("NO ENVIADA", "COSCO: el portal no aceptó el envío y marcó: «Gross Weight is required»."),
               ("EMITIDA", "MSC emitida con éxito (Booking: EBKG20000001)")]

    def estado(self, i):
        return self.ESTADOS[i % len(self.ESTADOS)]

    def test_consola_lee_lo_mismo(self):
        hoja = TestDiferencias.HOJA
        for nav, filas in CONSOLA.items():
            for i, (fila, *_) in enumerate(filas):
                e, det = self.estado(i)
                self.mod.escribir_estado(nav, fila, f"SOLICITUD {i + 1} · {e} · nave: X · {det}")
        wb = openpyxl.load_workbook(self.planilla)
        for nav, filas in CONSOLA.items():
            with self.subTest(naviera=nav):
                col = self.mod._columna_estado(wb[hoja[nav]])
                escritos = [str(wb[hoja[nav]].cell(fila, col).value) for fila, *_ in filas]
                self.assertEqual([t.split(" · ")[1] for t in escritos],
                                 [self.estado(i)[0] for i in range(len(filas))])         # control: quedaron escritos
                self.assertEqual(foto_consola(self.mod.leer_reservas(nav)), CONSOLA[nav])

    def test_web_lee_lo_mismo(self):
        for hoja, filas in WEB.items():
            resultados = {str(f[1]): dict(zip(("estado", "detalle"), self.estado(i))) for i, f in enumerate(filas)}
            ws = openpyxl.load_workbook(BytesIO(self.mod._planilla_con_resultados(str(self.planilla), hoja,
                                                                                  resultados)))[hoja]
            with self.subTest(hoja=hoja):
                self.assertEqual([str(ws.cell(f[1], 8).value).split(" · ")[0] for f in filas],
                                 [self.estado(i)[0] for i in range(len(filas))])         # control: quedaron escritos
                self.assertEqual(foto_web(self.mod.leer_filas(ws)[2]), WEB[hoja])


class TestDestinoFinalConGuion(ConPlanilla):
    def test_un_guion_en_el_destino_final_es_la_celda_vacia(self):
        # Decisión de Marcelo (encargo 43, CICLO-maersk-elige-exacto.md): «-» en el destino final, sin otra cosa, se
        # trata igual que esa celda vacía: los dos lectores toman el destino original y, sin él, la fila no trae destino
        # (_falta_en_la_fila). Hasta ahí, el destino era «-», y la fila quedaba NO ENVIADA con «a la fila le falta el
        # destino» (encargo 42; antes, las navieras buscaban «-»). Otro texto, como «--» o «N/A», queda como viene.
        wb = openpyxl.load_workbook(self.planilla)
        wb["MAERSK"]["D5"] = "-"                                    # con el puerto de descarga: el destino es ese
        wb["MAERSK"]["C6"], wb["MAERSK"]["D6"] = None, " - "        # sin él: la fila no trae destino
        wb["ONE"]["D5"] = "--"                                      # otro texto, como viene
        wb.save(self.planilla)
        m = self.mod
        libro = openpyxl.load_workbook(self.planilla, data_only=True)
        lectores = {"consola": {f["fila"]: f for f in m.leer_reservas("maersk")},
                    "web": {f["fila"]: f for f in m.leer_filas(libro["MAERSK"])[2]}}
        for lector, filas in lectores.items():
            with self.subTest(lector=lector):
                self.assertEqual([(filas[f]["destino_orig"], filas[f]["destino_final"]) for f in (5, 6)],
                                 [("PUERTO BETA", "PUERTO BETA"), ("", "")])
                self.assertEqual([m._falta_en_la_fila(filas[f]) for f in (5, 6)], ["", m.FILA_SIN_DESTINO])
        self.assertEqual(({f["fila"]: f for f in m.leer_reservas("one")}[5]["destino_final"],
                          {f["fila"]: f for f in m.leer_filas(libro["ONE"])[2]}[5]["destino_final"]), ("--", "--"))
        self.assertEqual([m._destino_final_de(x) for x in ("-", " - ", "--", "N/A", "", None, "CIUDAD GAMMA")],
                         ["", "", "--", "N/A", "", None, "CIUDAD GAMMA"])


if __name__ == "__main__":
    unittest.main()
