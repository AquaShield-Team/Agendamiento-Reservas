# -*- coding: utf-8 -*-
"""Panel web (lanzar_web) en un puerto propio, dentro de la sandbox.

Las navieras se reemplazan por login y reservador FALSOS y Playwright por uno
que solo registra: ninguna prueba abre un navegador ni toca un portal. Lo que
se fotografía es la orquestación: rutas, credenciales, planilla subida,
corridas por hoja y por consolidado, pausa, detener, modo login, el modo que manda la página al armar, y el
arranque (el puerto siguiente si otro programa tiene el base, sin cerrarlo; relevo de una instancia previa).
"""
import ast
import contextlib
import http.server
import io
import json
import os
import random
import re
import socket
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from io import BytesIO
from pathlib import Path
from unittest import mock
from urllib.parse import quote

import openpyxl

import sinteticos
import soporte

TESTS = Path(__file__).resolve().parent
RUTAS_JS = ["/api/apagar", "/api/config", "/api/continuar", "/api/correr", "/api/credenciales", "/api/descargar",
            "/api/detener", "/api/estado", "/api/filas", "/api/login", "/api/planilla"]


def planilla_bytes(sb):
    ruta = sinteticos.crear_planilla(sb / "_planilla_para_subir.xlsx")
    datos = ruta.read_bytes()
    ruta.unlink()
    return datos


Navieras = soporte.Navieras


class PanelBase(soporte.CasoAQ):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.srv = soporte.Servidor(cls.mod).iniciar()
        cls.xlsx = planilla_bytes(cls.sb)

    @classmethod
    def tearDownClass(cls):
        cls.srv.apagar()

    def subir(self, nombre="Prueba.xlsx"):
        return self.srv.json(f"/api/planilla?nombre={quote(nombre)}", self.xlsx)

    def correr(self, hoja, filas, usuario="op_prueba"):
        # «modo»: el que muestra la página, que en la sandbox es prueba (encargo 52).
        st, r = self.srv.json("/api/correr", {"hoja": hoja, "usuario": usuario, "filas": filas, "modo": False})
        self.assertEqual((st, r), (200, {"ok": True}))
        return self.srv.esperar_fin()


class TestPanelRutas(PanelBase):
    def test_recursos_del_panel(self):
        for ruta, tipo, marca in (("/", "text/html", b"<html"), ("/app.js", "application/javascript", b"fetch"),
                                  ("/estilo.css", "text/css", b"{"), ("/fuentes.css", "text/css", b"@font-face")):
            st, cab, cuerpo = self.srv.completo(ruta)
            with self.subTest(ruta=ruta):
                self.assertEqual(st, 200)
                self.assertTrue(cab["Content-Type"].startswith(tipo))
                self.assertIn(marca, cuerpo)
        self.assertEqual(self.srv.get("/api/no-existe"), (404, b"no existe"))
        self.assertEqual(self.srv.post("/api/no-existe"), (404, b"no existe"))

    def test_config_del_panel(self):
        st, cfg = self.srv.json("/api/config")
        self.assertEqual(st, 200)
        self.assertIs(cfg["modo_emision"], False)
        self.assertEqual(cfg["version"], self.mod.VERSION)
        self.assertEqual(cfg["usuarios"], ["op_prueba", "op_vacio"])
        self.assertEqual([(n["clave"], n["reserva"]) for n in cfg["navieras"]],
                         [(k, True) for k in sinteticos.NAVIERAS])

    def test_config_ausente_da_500(self):
        ruta = self.sb / "config.json"
        original = ruta.read_bytes()
        ruta.unlink()
        try:
            st, cuerpo = self.srv.get("/api/config")
        finally:
            ruta.write_bytes(original)
        self.assertEqual(st, 500)
        self.assertIn(b"No existe config.json (copia config.example.json", cuerpo)

    def test_llave_invalida_se_avisa_en_pantalla(self):
        ruta = self.sb / "config.json"
        original = ruta.read_bytes()
        cfg = json.loads(original)
        cfg["opciones"]["emitir_reservas"] = "true"            # «true» entre comillas: texto, no booleano
        soporte.escribir_json(ruta, cfg)
        try:
            _, antes = self.srv.json("/api/estado")
            st, config = self.srv.json("/api/config")
            _, despues = self.srv.json(f"/api/estado?desde={antes['cursor']}")
        finally:
            ruta.write_bytes(original)
        self.assertEqual((st, config["modo_emision"]), (200, False))
        self.assertEqual(len(despues["lineas"]), 1)
        self.assertIn('«emitir_reservas» vale "true" y solo se abre con true (sin comillas)', despues["lineas"][0])

    def test_rutas_del_js_existen_en_el_servidor(self):
        fuente = Path(self.mod.__file__).read_text(encoding="utf-8")
        js = set(re.findall(r"""['"`](/api/[a-z]+)""", self.mod.JS_INDEX))
        tree = ast.parse(fuente)
        f = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "lanzar_web")
        atendidas = set()
        for n in ast.walk(f):
            if isinstance(n, ast.Compare) and ast.unparse(n.left) == "u.path":
                for c in ast.walk(n.comparators[0]):
                    if isinstance(c, ast.Constant) and isinstance(c.value, str):
                        atendidas.add(c.value)
        self.assertEqual(sorted(js), RUTAS_JS)
        self.assertEqual(sorted(atendidas - {"/", "/index.html", "/app.js", "/estilo.css", "/fuentes.css"}),
                         sorted(RUTAS_JS + ["/api/salir", "/api/shutdown"]))


class TestPanelCredenciales(PanelBase):
    def tearDown(self):
        soporte.escribir_json(self.sb / "config.json", sinteticos.config())
        super().tearDown()

    def test_credenciales_nunca_devuelven_la_clave(self):
        st, cuerpo = self.srv.get("/api/credenciales?usuario=op_prueba")
        self.assertEqual(st, 200)
        self.assertNotIn(b"clave-falsa", cuerpo)
        r = json.loads(cuerpo)
        for n in r["navieras"]:
            self.assertEqual(set(n), {"clave", "nombre", "usuario", "contrato", "tiene_clave"})
            self.assertTrue(n["tiene_clave"])
            self.assertEqual(n["usuario"], f"usuario.{n['clave']}@ejemplo.invalid")
        _, vacio = self.srv.json("/api/credenciales?usuario=op_vacio")
        self.assertEqual({n["tiene_clave"] for n in vacio["navieras"]}, {False})

    def test_guardar_credenciales(self):
        leer = lambda: json.loads((self.sb / "config.json").read_text(encoding="utf-8"))["usuarios"]["op_nuevo"]
        st, r = self.srv.json("/api/credenciales", {"usuario": "op_nuevo", "descripcion": "Nuevo",
                                                    "navieras": {"one": {"usuario": "u1", "clave": "c1", "contrato": "k1"}}})
        self.assertEqual((st, r["ok"], r["operadores"]), (200, True, ["op_prueba", "op_vacio", "op_nuevo"]))
        op = leer()
        self.assertEqual(sorted(op), sorted(["descripcion"] + list(sinteticos.NAVIERAS)))   # nace con las seis
        self.assertEqual(op["one"], {"usuario": "u1", "clave": "c1", "contrato": "k1"})
        self.assertEqual(op["msc"], {"usuario": "", "clave": "", "contrato": ""})
        # clave ausente: se conserva la que había
        self.srv.json("/api/credenciales", {"usuario": "op_nuevo", "navieras": {"one": {"usuario": "u2", "contrato": "k1"}}})
        self.assertEqual(leer()["one"], {"usuario": "u2", "clave": "c1", "contrato": "k1"})
        # clave "" : se borra
        self.srv.json("/api/credenciales", {"usuario": "op_nuevo", "navieras": {"one": {"usuario": "u2", "clave": ""}}})
        self.assertEqual(leer()["one"]["clave"], "")
        self.assertFalse((self.sb / "config.json.tmp").exists())
        self.assertEqual(self.srv.json("/api/credenciales", {"navieras": {}}),
                         (400, {"error": "falta el nombre del operador"}))

    # --- El reemplazo de config.json que Windows niega un instante (CICLO-cola-cuatro-items.md) ---
    def negar_reemplazo(self, veces):
        """os.replace niega el reemplazo las primeras 'veces' veces, como Windows cuando otro proceso (el antivirus) tiene
        abierto el archivo un instante; después reemplaza de verdad. Devuelve el parche y la lista de intentos."""
        real, intentos = self.mod._os.replace, []

        def replace(origen, destino):
            intentos.append(Path(destino).name)
            if len(intentos) <= veces:
                raise PermissionError(13, "Acceso denegado (sintético)")
            return real(origen, destino)
        return mock.patch.object(self.mod._os, "replace", replace), intentos

    def test_guardar_reintenta_si_windows_niega_el_reemplazo(self):
        # Medido: 47 de 900 guardados seguidos respondían 500 «[WinError 5] Acceso denegado» al reemplazar config.json,
        # y hasta 45f2171 test_guardar_credenciales caía así en 1 de 12 suites completas.
        parche, intentos = self.negar_reemplazo(3)
        with parche:
            st, r = self.srv.json("/api/credenciales", {"usuario": "op_nuevo",
                                                        "navieras": {"one": {"usuario": "u9", "clave": "c9"}}})
        self.assertEqual((st, r["ok"]), (200, True))
        self.assertEqual(intentos, ["config.json"] * 4)
        cfg = json.loads((self.sb / "config.json").read_text(encoding="utf-8"))
        self.assertEqual(cfg["usuarios"]["op_nuevo"]["one"], {"usuario": "u9", "clave": "c9", "contrato": ""})
        self.assertFalse((self.sb / "config.json.tmp").exists())

    def test_guardar_se_rinde_si_el_reemplazo_sigue_negado(self):
        antes = (self.sb / "config.json").read_bytes()
        self.addCleanup(lambda: (self.sb / "config.json.tmp").unlink(missing_ok=True))
        parche, intentos = self.negar_reemplazo(10 ** 6)
        with parche:
            st, cuerpo = self.srv.pedir("/api/credenciales", {"usuario": "op_nuevo", "navieras": {}})
        self.assertEqual(st, 500)
        self.assertIn("Acceso denegado", cuerpo.decode("utf-8"))            # el panel muestra el error
        self.assertEqual(len(intentos), 20)
        self.assertEqual((self.sb / "config.json").read_bytes(), antes)     # config.json queda como estaba

    def test_reemplazar_solo_reintenta_el_acceso_denegado(self):
        llamadas = []

        def replace(origen, destino):
            llamadas.append(destino)
            raise FileNotFoundError(2, "no está (sintético)")
        with mock.patch.object(self.mod._os, "replace", replace):
            with self.assertRaises(FileNotFoundError):
                self.mod._reemplazar("a.tmp", "a", pausa=0)
        self.assertEqual(llamadas, ["a"])                                   # otro error sube de inmediato

    # --- Los dos paneles escriben config.json con el mismo ayudante atómico (CICLO-cola-siete-items.md) ---
    def test_escribir_config_es_atomica_y_reintenta(self):
        ruta = self.sb / "config.json"
        cfg = {"usuarios": {"op_ñandú": {"descripcion": "Operación de prueba"}}}
        parche, intentos = self.negar_reemplazo(2)
        with parche:
            self.assertEqual(self.mod._escribir_config(cfg, ruta), 2)                   # reintentó dos veces
        self.assertEqual(intentos, ["config.json"] * 3)
        # El mismo formato que escribían los dos paneles: UTF-8 sin escapar, con sangría de 2.
        self.assertEqual(ruta.read_text(encoding="utf-8"), json.dumps(cfg, ensure_ascii=False, indent=2))
        self.assertFalse((self.sb / "config.json.tmp").exists())

    def test_escribir_config_no_toca_el_archivo_si_el_reemplazo_sigue_negado(self):
        ruta = self.sb / "config.json"
        antes = ruta.read_bytes()
        self.addCleanup(lambda: (self.sb / "config.json.tmp").unlink(missing_ok=True))
        parche, intentos = self.negar_reemplazo(10 ** 6)
        with parche, self.assertRaises(PermissionError):
            self.mod._escribir_config({"usuarios": {}}, ruta)
        self.assertEqual(len(intentos), 20)
        self.assertEqual(ruta.read_bytes(), antes)                                     # ni truncado ni a medias

    def test_los_dos_paneles_escriben_config_con_el_mismo_ayudante(self):
        # Hasta 0af2d5b, el panel de escritorio abría config.json con open(w), que lo trunca al abrir: un fallo a mitad
        # lo dejaba vacío, y con él las credenciales de los portales.
        src = Path(self.mod.__file__).read_text(encoding="utf-8")
        arbol = ast.parse(src)
        padres = {}
        for n in ast.walk(arbol):
            for h in ast.iter_child_nodes(n):
                padres[h] = n

        def funcion(n):
            while n in padres and not isinstance(n, ast.FunctionDef):
                n = padres[n]
            return n.name
        llamadas = [c for c in ast.walk(arbol) if isinstance(c, ast.Call) and getattr(c.func, "id", "") == "_escribir_config"]
        self.assertEqual(sorted((funcion(c), ast.unparse(c)) for c in llamadas),
                         [("do_POST", "_escribir_config(cfg, BASE / 'config.json')"),
                          ("guardar", "_escribir_config(self.cfg, BASE / 'config.json')")])
        # Nadie más abre config.json para escribir, ni escribe el temporal por su cuenta.
        escrituras = [ast.unparse(c) for c in ast.walk(arbol) if isinstance(c, ast.Call)
                      and getattr(c.func, "id", "") == "open" and len(c.args) > 1
                      and isinstance(c.args[1], ast.Constant) and "w" in str(c.args[1].value)]
        self.assertEqual([e for e in escrituras if "config" in e or e.startswith("open(tmp")],
                         ["open(tmp, 'w', encoding='utf-8')"])
        self.assertEqual(src.count('".json.tmp"'), 1)


class TestPanelPlanilla(PanelBase):
    def test_planilla_subida_va_al_temporal_con_nombre_fijo(self):
        st, r = self.subir()
        self.assertEqual(st, 200)
        self.assertEqual(r["nombre"], "Prueba.xlsx")
        self.assertEqual([(h["hoja"], h["filas"], h["reservador"]) for h in r["hojas"]],
                         [("CONSOLIDADO", 7, True), ("ONE", 5, True), ("MSC", 3, True), ("CMA-CGM", 3, True),
                          ("COSCO", 2, True), ("HYUNDAI", 2, True), ("MAERSK", 2, True)])
        destino = Path(tempfile.gettempdir()) / "aquashield_planilla.xlsx"
        self.assertTrue(self.sb.resolve() in destino.resolve().parents, "el temporal no está en la sandbox")
        self.assertEqual(destino.read_bytes(), self.xlsx)

    def test_filas_de_una_hoja(self):
        self.subir()
        _, r = self.srv.json("/api/filas?hoja=ONE")
        self.assertEqual([f["fila"] for f in r["filas"]], [5, 6, 7, 9, 10])
        _, r = self.srv.json("/api/filas?hoja=NO-EXISTE")
        self.assertEqual(r, {"filas": []})

    def test_descarga_de_resultados(self):
        self.subir("Mi Planilla.xlsx")
        temporal = Path(tempfile.gettempdir()) / "aquashield_planilla.xlsx"
        antes = temporal.read_bytes()
        with self.mod._LOCK:
            self.mod._WEB.update({"hoja": "ONE", "resultados": {"5": {"estado": "EMITIDA", "booking": "BKGFALSO03"}}})
        st, cab, cuerpo = self.srv.completo("/api/descargar")
        self.assertEqual(st, 200)
        self.assertRegex(cab["Content-Disposition"],
                         r'^attachment; filename="Mi Planilla - resultados \d\d-\d\d_\d{4}\.xlsx"$')
        ws = openpyxl.load_workbook(BytesIO(cuerpo))["ONE"]
        self.assertEqual((ws["H5"].value, ws["I5"].value), ("EMITIDA", "BKGFALSO03"))
        self.assertEqual(temporal.read_bytes(), antes, "la descarga tocó la planilla subida")

    def test_numero_de_reserva_igual_en_el_panel_y_en_la_planilla(self):
        # La columna «N° de reserva emitida» dice lo mismo en el panel (n_reserva de /api/estado) y en la planilla: el
        # número de la EMITIDA; «lista para emitir» en la OK-EJEMPLO, que no se emitió; nada en las demás. Hasta el
        # encargo 34, el panel decía «✓ Confirmada» en la OK-EJEMPLO y la planilla dejaba la celda vacía
        # (CICLO-one-y-maersk-eligen-la-nave.md).
        self.subir()
        resultados = {"5": {"estado": "EMITIDA", "detalle": "ONE emitida con éxito (Booking: BKGFALSO11)", "booking": ""},
                      "6": {"estado": "OK-EJEMPLO", "detalle": "armada hasta Review; SIN emitir", "booking": ""},
                      "7": {"estado": "NO ENVIADA", "detalle": "ONE: no elegí salida", "booking": ""}}
        with self.mod._LOCK:
            self.mod._WEB.update({"hoja": "ONE", "resultados": resultados})
        try:
            _, _, cuerpo = self.srv.completo("/api/descargar")
            _, estado = self.srv.json("/api/estado")
        finally:
            with self.mod._LOCK:
                self.mod._WEB.update({"resultados": {}})
        ws = openpyxl.load_workbook(BytesIO(cuerpo))["ONE"]
        planilla = {f: ws[f"I{f}"].value or "" for f in resultados}
        panel = {f: estado["resultados"][f]["n_reserva"] for f in resultados}
        self.assertEqual(planilla, {"5": "BKGFALSO11", "6": "lista para emitir", "7": ""})
        self.assertEqual(panel, planilla)
        self.assertEqual(self.mod.LISTA_PARA_EMITIR, "lista para emitir")


    def test_descarga_avisa_en_pantalla_y_en_el_log(self):
        # De punta a punta: lo que la descarga no pudo escribir (la fila 1 es un título combinado)
        # sale en el registro que muestra el panel y en el log.txt de la corrida.
        self.subir()
        carpeta = self.sb / "logs" / "web_one_prueba_descarga"
        _, antes = self.srv.json("/api/estado")
        with self.mod._LOCK:
            self.mod._WEB.update({"hoja": "ONE", "carpeta": str(carpeta),
                                  "resultados": {"1": {"estado": "EMITIDA", "detalle": "x", "booking": "BKGFALSO07"}}})
        try:
            st, _, _ = self.srv.completo("/api/descargar")
            _, despues = self.srv.json(f"/api/estado?desde={antes['cursor']}")
        finally:
            with self.mod._LOCK:
                self.mod._WEB.update({"carpeta": "", "resultados": {}})
        self.assertEqual(st, 200)
        log = (carpeta / "log.txt").read_text(encoding="utf-8")
        for texto in ("· ⚠ No pude escribir en la planilla descargada el resultado de la fila 1: error inesperado "
                      "(AttributeError: ", "Anótalo a mano: EMITIDA · BKGFALSO07 · x",
                      "· ⚠ 1 resultado(s) no quedaron en la planilla descargada (filas 1); cada uno está arriba en este log."):
            with self.subTest(texto=texto[:40]):
                self.assertIn(texto, log)
                self.assertTrue(any(texto in l for l in despues["lineas"]), despues["lineas"])


class TestPanelSinPlanilla(PanelBase):
    def test_descarga_sin_planilla(self):
        self.assertEqual(self.srv.get("/api/descargar"), (400, "No hay planilla cargada.".encode("utf-8")))
        self.assertEqual(self.srv.json("/api/filas?hoja=ONE"), (200, {"filas": []}))


class TestPanelCorridas(PanelBase):
    def setUp(self):
        super().setUp()
        self.subir()

    def test_correr_rechazos(self):
        self.assertEqual(self.srv.json("/api/correr", {"hoja": "OTRA", "filas": [5], "modo": False}),
                         (400, {"error": "'OTRA' no tiene automatización"}))
        with self.mod._LOCK:
            self.mod._WEB["corriendo"] = True
        try:
            self.assertEqual(self.srv.json("/api/correr", {"hoja": "ONE", "filas": [5], "modo": False}),
                             (409, {"error": "ya hay una corrida en curso"}))
        finally:
            with self.mod._LOCK:
                self.mod._WEB["corriendo"] = False

    def test_no_arma_si_la_pagina_muestra_otro_modo(self):
        # Decisión de Marcelo, encargo 52 (CICLO-puerto-libre-y-modo-al-armar.md): antes de armar, el panel vuelve a
        # leer su modo, y si no es el que la página dice mostrar («modo», en /api/correr), no corre y avisa que hay que
        # recargarla. La página lo manda desde que pinta su aviso de modo (MODO_PAGINA: test_envio); la que no lo
        # manda, como una pestaña de antes de este encargo, tampoco arma. Nada corre: ni el navegador ni el estado del
        # panel.
        texto = ("No armé las reservas: este panel está en modo {}, y la página {}. Recarga la página (F5) para ver el "
                 "modo del panel, y vuelve a armarlas.")
        prueba, emision = "prueba (sin emitir)", "EMISIÓN (reservas reales)"
        sin_decir = texto.format(prueba, "no dice qué modo muestra")
        casos = [({"modo": True}, texto.format(prueba, f"muestra el modo {emision}")), ({}, sin_decir),
                 ({"modo": None}, sin_decir), ({"modo": "false"}, sin_decir), ({"modo": 0}, sin_decir)]
        pedido = {"hoja": "ONE", "usuario": "op_prueba", "filas": [5]}
        with Navieras(self.mod) as n:
            for extra, error in casos:
                with self.subTest(extra=extra):
                    self.assertEqual(self.srv.json("/api/correr", {**pedido, **extra}), (409, {"error": error}))
            # Con una llave del candado abierta, el panel está en modo emisión: la página que muestra el de prueba
            # tampoco arma. (Con «modo» true no se prueba: armaría las reservas en modo emisión.)
            try:
                os.environ[soporte.LLAVE_ENTORNO] = "1"
                respuesta = self.srv.json("/api/correr", {**pedido, "modo": False})
            finally:
                soporte.apagar_llaves()
        self.assertEqual(respuesta, (409, {"error": texto.format(emision, f"muestra el modo {prueba}")}))
        self.assertEqual((self.mod._WEB["corriendo"], n.pw.lanzamientos, n.reservas), (False, [], []))

    def test_corrida_de_una_hoja(self):
        # Con las filas 5 y 6: hasta a198645 eran la 5 y la 7, que no trae nave y ahora no va al portal
        # (test_fila_sin_nave_no_abre_el_portal, CICLO-solo-la-nave-pedida.md).
        def respuesta(nav, rsv, on_pausa):
            if rsv["fila"] == 6:
                return ("EMITIDA", "ONE emitida con éxito (Booking: FALSO12345)")
            return ("OK-EJEMPLO", "armada falsa")
        antes = {p.name for p in (self.sb / "logs").glob("*")}
        with Navieras(self.mod, respuesta=respuesta) as n:
            e = self.correr("ONE", [5, 6])
        self.assertEqual(n.logins, [("one", "usuario.one@ejemplo.invalid")])
        self.assertEqual(n.reservas, [("one", 5, "usuario.one@ejemplo.invalid"), ("one", 6, "usuario.one@ejemplo.invalid")])
        self.assertEqual(e["estado"], "Terminado.")
        self.assertEqual({k: (v["estado"], v["booking"], v["nave"]) for k, v in e["resultados"].items()},
                         {"5": ("OK-EJEMPLO", "", "NAVE PRUEBA UNO"), "6": ("EMITIDA", "FALSO12345", "NAVE PRUEBA DOS")})
        lanz = n.pw.lanzamientos
        self.assertEqual(len(lanz), 1)
        self.assertEqual(Path(lanz[0]["user_data_dir"]), self.sb / "perfiles" / "op_prueba")
        self.assertEqual((lanz[0]["channel"], lanz[0]["headless"], lanz[0]["no_viewport"], lanz[0]["ignore_default_args"]),
                         ("chrome", False, True, ["--enable-automation"]))
        self.assertIs(lanz[0]["chromium_sandbox"], True)          # con el sandbox de Chrome (encargo 53)
        self.assertEqual(lanz[0]["args"], self.mod._args_chrome())
        self.assertTrue(n.pw.contextos[0].cerrado)
        self.assertEqual(n.pw.contextos[0].scripts, [self.mod.STEALTH_JS])
        self.assertIn("INICIO · ONE · operador op_prueba · 2 reserva(s)", e["lineas"][0])
        carpeta = Path(self.mod._WEB["carpeta"])
        self.assertEqual(carpeta.parent, self.sb / "logs")
        # Con «_2», «_3»… si otra corrida ya tomó ese segundo; la carpeta tiene que ser nueva (hasta 393f504
        # compartía la de la prueba anterior 4 de cada 5 veces, y esta prueba no lo veía).
        self.assertRegex(carpeta.name, r"^web_one_op_prueba_\d{8}_\d{6}(_\d+)?$")
        self.assertNotIn(carpeta.name, antes)
        self.assertTrue((carpeta / "log.txt").exists())

    def test_dice_su_candado_al_empezar(self):
        # Cada corrida del panel dice al empezar, en su registro y en log.txt, con qué candado corre y por cuál llave
        # (_anotar_candado; decisión de Marcelo, encargo 51). Con las tres apagadas, modo prueba; con llaves, cuáles
        # (aquí, dichas por _llaves_abiertas solo para la línea: el candado de la guarda sigue cerrado).
        with Navieras(self.mod):
            e = self.correr("ONE", [5])
        self.assertIn("· INICIO · ONE · operador op_prueba · 1 reserva(s)", e["lineas"][0])
        self.assertIn("· 🛡️ Candado de emisión cerrado: modo prueba, sin ninguna de sus tres llaves abierta",
                      e["lineas"][1])
        log = (Path(self.mod._WEB["carpeta"]) / "log.txt").read_text(encoding="utf-8").splitlines()
        self.assertEqual(log[:2], e["lineas"][:2])
        llaves = (lambda todas=False: ["la primera", "la segunda"] if todas else [])
        with Navieras(self.mod), mock.patch.object(self.mod, "_llaves_abiertas", llaves):
            e = self.correr("ONE", [5])
        self.assertIn("· 🔴 Candado de emisión abierto: modo EMISIÓN, por la primera y la segunda. Cada reserva que "
                      "llegue al botón final se envía a la naviera.", e["lineas"][1])

    def test_dos_corridas_en_el_mismo_segundo(self):
        # Con el reloj detenido, dos corridas seguidas caen en el mismo segundo: cada una con su carpeta
        # y su log (CICLO-carpetas-unicas.md).
        carpetas = []
        with soporte.reloj_detenido(self.mod), Navieras(self.mod):
            for _ in range(2):
                self.correr("ONE", [5])
                carpetas.append(Path(self.mod._WEB["carpeta"]))
        self.assertEqual([c.name for c in carpetas],
                         ["web_one_op_prueba_20260102_030405", "web_one_op_prueba_20260102_030405_2"])
        self.assertEqual([len(re.findall(r"· INICIO", (c / "log.txt").read_text(encoding="utf-8")))
                          for c in carpetas], [1, 1])

    def test_corrida_con_estados_tras_el_envio(self):
        # De punta a punta: lo que devuelve el reservador llega al panel y a la planilla descargada.
        # El número va solo con EMITIDA, aunque el error del portal cite un código. La 7 no trae nave: su NO ENVIADA
        # lo pone el panel sin ir al portal. Hasta a198645, la 5 volvía NO ENVIADA del reservador y la 7, EMITIDA
        # (CICLO-solo-la-nave-pedida.md).
        respuestas = {
            5: ("EMITIDA", "ONE emitida con éxito (Booking: SCLG10000003)"),
            6: ("ENVIADA – REVISAR EN PORTAL", "ONE: se pulsó «Submit» y el portal mostró un error: «Booking: "
                                               "COSFALSO9876 ya existe». Revisa en el portal.")}
        with Navieras(self.mod, respuesta=lambda nav, rsv, on_pausa: respuestas[rsv["fila"]]):
            e = self.correr("ONE", [5, 6, 7])
        self.assertEqual({k: (v["estado"], v["booking"]) for k, v in e["resultados"].items()},
                         {"5": ("EMITIDA", "SCLG10000003"), "6": ("ENVIADA – REVISAR EN PORTAL", ""), "7": ("NO ENVIADA", "")})
        st, _, cuerpo = self.srv.completo("/api/descargar")
        self.assertEqual(st, 200)
        ws = openpyxl.load_workbook(BytesIO(cuerpo))["ONE"]
        self.assertEqual([(ws[f"H{f}"].value, ws[f"I{f}"].value) for f in (5, 6, 7)],
                         [("EMITIDA", "SCLG10000003"),
                          (f"ENVIADA – REVISAR EN PORTAL · {respuestas[6][1]}", None),
                          ("NO ENVIADA · la fila no trae una nave para reservar", None)])

    def test_corrida_sin_filas_no_abre_el_navegador(self):
        # Hasta 42247d4, con una sola naviera, abría el navegador y entraba al portal igual
        # (CICLO-cola-nueve-items.md).
        aviso = ("⚠ No hay filas elegidas: no abro el navegador ni entro a ningún portal. Elige al menos una fila y "
                 "vuelve a armar las reservas.")
        for hoja in ("ONE", "CONSOLIDADO"):
            with self.subTest(hoja=hoja), Navieras(self.mod) as n:
                e = self.correr(hoja, [])
                self.assertEqual((len(n.pw.lanzamientos), len(n.logins), n.reservas), (0, 0, []))
                self.assertEqual((e["estado"], e["resultados"]), ("Terminado.", {}))
                self.assertEqual(sum(aviso in l for l in e["lineas"]), 1)
                # Ni el aviso de las filas sin nave, que es para una naviera con filas (desde el encargo 28).
                self.assertFalse(any("Ninguna fila de" in l for l in e["lineas"]))

    def test_fila_sin_nave_no_abre_el_portal(self):
        # Una fila sin nave (vacía o con un comodín) queda NO ENVIADA con el motivo, en pantalla, en log.txt y en la
        # planilla, y no se abre el portal por ella (decisión de Marcelo, CICLO-solo-la-nave-pedida.md). Hasta a198645
        # iba al reservador: ONE y COSCO devolvían OMITIDO solo con la celda vacía, y con un comodín entraban al portal.
        motivo = "la fila no trae una nave para reservar"
        self.subir_con({("ONE", 6): "Primera disponible"})     # un comodín; la 7 viene vacía y la 10, «POR COMPLETAR»
        with Navieras(self.mod) as n:
            e = self.correr("ONE", [5, 6, 7, 10])
        self.assertEqual((len(n.pw.lanzamientos), n.reservas), (1, [("one", 5, "usuario.one@ejemplo.invalid")]))
        self.assertEqual({k: (v["estado"], v["detalle"], v["booking"]) for k, v in e["resultados"].items() if k != "5"},
                         {f: ("NO ENVIADA", motivo, "") for f in ("6", "7", "10")})
        log = (Path(self.mod._WEB["carpeta"]) / "log.txt").read_text(encoding="utf-8")
        for f in (6, 7, 10):
            linea = (f"✗ fila {f}: NO ENVIADA · {motivo}. No entro al portal por ella: escribe la nave en la planilla y "
                     "vuelve a armar la reserva.")
            with self.subTest(fila=f):
                self.assertIn(linea, log)
                self.assertEqual(sum(linea in l for l in e["lineas"]), 1)
        st, _, cuerpo = self.srv.completo("/api/descargar")
        self.assertEqual(st, 200)
        ws = openpyxl.load_workbook(BytesIO(cuerpo))["ONE"]
        self.assertEqual([(ws[f"H{f}"].value, ws[f"I{f}"].value) for f in (6, 7, 10)],
                         [(f"NO ENVIADA · {motivo}", None)] * 3)
        # Si ninguna fila de la naviera trae nave, no se abre el navegador ni se entra al portal.
        aviso = "Ninguna fila de ONE trae una nave para reservar: no abro el navegador ni entro al portal."
        with Navieras(self.mod) as n:
            e = self.correr("ONE", [6, 7])
        self.assertEqual((len(n.pw.lanzamientos), n.logins, n.reservas), (0, [], []))
        self.assertEqual(sorted(e["resultados"], key=int), ["6", "7"])
        self.assertEqual(sum(aviso in l for l in e["lineas"]), 1)

    def test_fila_con_solo_el_prefijo_no_abre_el_portal(self):
        # Una fila con solo el prefijo de la naviera tampoco pide una nave: /api/filas lo dice, y la corrida la deja NO
        # ENVIADA sin abrir el navegador (decisión de Marcelo, CICLO-cola-seis-items.md). Hasta f68ce6c iba al portal, y
        # el reservador buscaba «ONE» en las tarjetas.
        motivo = "la fila no trae una nave para reservar"
        self.subir_con({("ONE", 5): "ONE"})
        st, d = self.srv.json("/api/filas?hoja=ONE")
        self.assertEqual((st, {f["fila"]: f["sin_nave"] for f in d["filas"]}[5]), (200, True))
        with Navieras(self.mod) as n:
            e = self.correr("ONE", [5])
        self.assertEqual((len(n.pw.lanzamientos), n.logins, n.reservas), (0, [], []))
        self.assertEqual((e["resultados"]["5"]["estado"], e["resultados"]["5"]["detalle"]), ("NO ENVIADA", motivo))

    def test_fila_manual_corrida_igual_queda_no_enviada(self):
        # /api/filas marca la fila MANUAL (el panel la deja sin marcar) y no la cuenta como sin nave. Si se corre igual,
        # queda NO ENVIADA con su motivo, en pantalla, en log.txt y en la planilla, sin entrar al portal por ella
        # (decisión de Marcelo, CICLO-cola-seis-items.md). Hasta f68ce6c iba al portal y buscaba «MANUAL» como nave.
        motivo = "la fila está marcada para hacerse a mano"
        self.subir_con({("ONE", 6): "Reserva MANUAL"})
        st, d = self.srv.json("/api/filas?hoja=ONE")
        self.assertEqual(st, 200)
        self.assertEqual({f["fila"]: (f["sin_nave"], f["manual"]) for f in d["filas"]},
                         {5: (False, False), 6: (False, True), 7: (True, False), 9: (True, False), 10: (True, False)})
        with Navieras(self.mod) as n:
            e = self.correr("ONE", [5, 6])
        self.assertEqual((len(n.pw.lanzamientos), n.reservas), (1, [("one", 5, "usuario.one@ejemplo.invalid")]))
        self.assertEqual((e["resultados"]["6"]["estado"], e["resultados"]["6"]["detalle"]), ("NO ENVIADA", motivo))
        linea = (f"✗ fila 6: NO ENVIADA · {motivo}. No entro al portal por ella: hazla a mano en el portal, o escribe "
                 "en la celda solo la nave, sin MANUAL, si quieres que la arme yo.")
        self.assertIn(linea, (Path(self.mod._WEB["carpeta"]) / "log.txt").read_text(encoding="utf-8"))
        self.assertEqual(sum(linea in l for l in e["lineas"]), 1)
        st, _, cuerpo = self.srv.completo("/api/descargar")
        ws = openpyxl.load_workbook(BytesIO(cuerpo))["ONE"]
        self.assertEqual((st, ws["H6"].value, ws["I6"].value), (200, f"NO ENVIADA · {motivo}", None))
        # Sola, no abre el navegador.
        with Navieras(self.mod) as n:
            self.correr("ONE", [6])
        self.assertEqual((len(n.pw.lanzamientos), n.logins, n.reservas), (0, [], []))

    def subir_con(self, naves):
        """Sube la planilla sintética con estas naves: {(hoja, fila): nave}."""
        wb = openpyxl.load_workbook(BytesIO(self.xlsx))
        for (hoja, fila), nave in naves.items():
            wb[hoja][f"E{fila}"] = nave
        cuerpo = BytesIO()
        wb.save(cuerpo)
        self.assertEqual(self.srv.json("/api/planilla?nombre=Naves.xlsx", cuerpo.getvalue())[0], 200)

    def test_fila_sin_nave_en_consolidado(self):
        # En CONSOLIDADO, cada fila sin nave queda NO ENVIADA una sola vez, con su naviera, y una naviera sin filas con
        # nave no abre su navegador (CICLO-solo-la-nave-pedida.md).
        self.subir_con({("CONSOLIDADO", 6): None, ("CONSOLIDADO", 7): "Cualquiera"})     # la de MSC y la de CMA
        with Navieras(self.mod) as n:
            e = self.correr("CONSOLIDADO", [5, 6, 7, 8])
        self.assertEqual([r[:2] for r in n.reservas], [("one", 5), ("cosco", 8)])
        self.assertEqual([l[0] for l in n.logins], ["one", "cosco"])
        self.assertEqual({k: v["estado"] for k, v in e["resultados"].items()},
                         {"5": "OK-EJEMPLO", "6": "NO ENVIADA", "7": "NO ENVIADA", "8": "OK-EJEMPLO"})
        for texto in ("✗ fila 6: NO ENVIADA", "✗ fila 7: NO ENVIADA",
                      "Ninguna fila de MSC trae una nave para reservar", "Ninguna fila de CMA-CGM trae una nave para reservar"):
            with self.subTest(texto=texto):
                self.assertEqual(sum(texto in l for l in e["lineas"]), 1)

    def subir_celdas(self, celdas):
        """Sube la planilla sintética con estas celdas: {(hoja, celda): valor}."""
        wb = openpyxl.load_workbook(BytesIO(self.xlsx))
        for (hoja, celda), valor in celdas.items():
            wb[hoja][celda] = valor
        cuerpo = BytesIO()
        wb.save(cuerpo)
        self.assertEqual(self.srv.json("/api/planilla?nombre=Ruta.xlsx", cuerpo.getvalue())[0], 200)

    def test_fila_sin_puerto_o_sin_destino_no_abre_el_portal(self):
        # Una fila con nave y sin puerto de carga o sin destino queda NO ENVIADA con el motivo, en pantalla, en log.txt
        # y en la planilla, sin abrir el portal por ella, como la fila sin nave (decisión de Marcelo, encargo 42,
        # CICLO-maersk-y-fila-sin-ruta.md). La 6, con solo el país en el puerto de carga; la 7, con nave y sin destino.
        puerto, destino = "a la fila le falta el puerto de carga", "a la fila le falta el destino"
        self.subir_celdas({("ONE", "B6"): ", CHILE", ("ONE", "E7"): "NAVE PRUEBA QUINCE", ("ONE", "C7"): None,
                           ("ONE", "D7"): None})
        with Navieras(self.mod) as n:
            e = self.correr("ONE", [5, 6, 7])
        self.assertEqual((len(n.pw.lanzamientos), n.reservas), (1, [("one", 5, "usuario.one@ejemplo.invalid")]))
        self.assertEqual({k: (v["estado"], v["detalle"], v["booking"]) for k, v in e["resultados"].items() if k != "5"},
                         {"6": ("NO ENVIADA", puerto, ""), "7": ("NO ENVIADA", destino, "")})
        log = (Path(self.mod._WEB["carpeta"]) / "log.txt").read_text(encoding="utf-8")
        for f, falta, que in ((6, puerto, "el puerto de carga"), (7, destino, "el destino")):
            linea = (f"✗ fila {f}: NO ENVIADA · {falta}. No entro al portal por ella: escribe {que} en la planilla y "
                     "vuelve a armar la reserva.")
            with self.subTest(fila=f):
                self.assertIn(linea, log)
                self.assertEqual(sum(linea in l for l in e["lineas"]), 1)
        st, _, cuerpo = self.srv.completo("/api/descargar")
        ws = openpyxl.load_workbook(BytesIO(cuerpo))["ONE"]
        self.assertEqual((st, [(ws[f"H{f}"].value, ws[f"I{f}"].value) for f in (6, 7)]),
                         (200, [(f"NO ENVIADA · {puerto}", None), (f"NO ENVIADA · {destino}", None)]))
        # Si ninguna fila con nave trae la ruta, no se abre el navegador ni se entra al portal.
        aviso = ("Ninguna fila de ONE con nave trae el puerto de carga y el destino: no abro el navegador ni entro al "
                 "portal.")
        with Navieras(self.mod) as n:
            e = self.correr("ONE", [6, 7])
        self.assertEqual((len(n.pw.lanzamientos), n.logins, n.reservas), (0, [], []))
        self.assertEqual(sorted(e["resultados"], key=int), ["6", "7"])
        self.assertEqual(sum(aviso in l for l in e["lineas"]), 1)
        # Sin nave y sin destino, el motivo es el de la nave (la 7 viene sin nave).
        self.subir_celdas({("ONE", "C7"): None, ("ONE", "D7"): None})
        with Navieras(self.mod) as n:
            e = self.correr("ONE", [7])
        self.assertEqual((e["resultados"]["7"]["detalle"], n.reservas), ("la fila no trae una nave para reservar", []))

    def test_fila_sin_puerto_en_consolidado(self):
        # En CONSOLIDADO, cada fila con su naviera: la de MSC sin puerto de carga queda NO ENVIADA una sola vez, y MSC
        # no abre su navegador (encargo 42).
        self.subir_celdas({("CONSOLIDADO", "B6"): None})
        with Navieras(self.mod) as n:
            e = self.correr("CONSOLIDADO", [5, 6, 7])
        self.assertEqual([r[:2] for r in n.reservas], [("one", 5), ("cma", 7)])
        self.assertEqual([l[0] for l in n.logins], ["one", "cma"])
        self.assertEqual({k: v["estado"] for k, v in e["resultados"].items()},
                         {"5": "OK-EJEMPLO", "6": "NO ENVIADA", "7": "OK-EJEMPLO"})
        for texto in ("✗ fila 6: NO ENVIADA · a la fila le falta el puerto de carga",
                      "Ninguna fila de MSC con nave trae el puerto de carga y el destino"):
            with self.subTest(texto=texto):
                self.assertEqual(sum(texto in l for l in e["lineas"]), 1)

    def test_filas_dicen_si_traen_nave(self):
        # /api/filas dice, con la regla del programa, qué fila no trae una nave para reservar: el panel la muestra y la
        # marca con eso (test_envio.TestLecturaDelPanel). Hasta a198645 el panel solo miraba la celda vacía.
        self.subir_con({("ONE", 6): "Primera disponible"})
        st, d = self.srv.json("/api/filas?hoja=ONE")
        self.assertEqual(st, 200)
        self.assertEqual({f["fila"]: f["sin_nave"] for f in d["filas"]}, {5: False, 6: True, 7: True, 9: True, 10: True})

    def test_filas_dicen_si_traen_la_ruta(self):
        # /api/filas dice, con la regla del programa (_falta_en_la_fila), qué fila con nave no trae su puerto de carga o
        # su destino: el panel la marca y la cuenta aparte antes de correr, como la fila sin nave (decisión de Marcelo,
        # encargo 43, CICLO-maersk-elige-exacto.md; test_envio.TestLecturaDelPanel). La fila sin nave, o MANUAL, no
        # lleva esas marcas: su motivo es el de la nave. La 6, con solo el país en el puerto de carga; la 7, con nave y
        # sin destino; la 9, sin nave y sin destino; la 10, MANUAL y sin puerto de carga.
        self.subir_celdas({("ONE", "B6"): ", CHILE", ("ONE", "E7"): "NAVE PRUEBA QUINCE", ("ONE", "C7"): None,
                           ("ONE", "D7"): None, ("ONE", "E10"): "Reserva MANUAL", ("ONE", "B10"): None})
        st, d = self.srv.json("/api/filas?hoja=ONE")
        self.assertEqual(st, 200)
        marcas = {f["fila"]: (f["sin_puerto"], f["sin_destino"], f["sin_nave"], f["manual"]) for f in d["filas"]}
        self.assertEqual(marcas, {5: (False, False, False, False), 6: (True, False, False, False),
                                  7: (False, True, False, False), 9: (False, False, True, False),
                                  10: (False, False, False, True)})

    def test_consolidado_reparte_por_naviera(self):
        with Navieras(self.mod) as n:
            e = self.correr("CONSOLIDADO", [5, 6, 7, 8, 9, 10, 11])
        self.assertEqual([r[:2] for r in n.reservas],
                         [("one", 5), ("msc", 6), ("cma", 7), ("cosco", 8), ("hyundai", 9), ("maersk", 10)])
        self.assertEqual([l[0] for l in n.logins], ["one", "msc", "cma", "cosco", "hyundai", "maersk"])
        self.assertEqual(len(n.pw.lanzamientos), 6)                  # un navegador por naviera
        self.assertNotIn("11", e["resultados"])                      # «OTRA LINEA» no se procesa
        self.assertEqual(sorted(e["resultados"], key=int), ["5", "6", "7", "8", "9", "10"])

    def test_login_fallido_no_reserva(self):
        with Navieras(self.mod, login_ok=False) as n:
            e = self.correr("ONE", [5, 7])
        # La 7 no trae nave: queda NO ENVIADA antes del login, y no depende de él. Hasta a198645 no tenía resultado
        # (CICLO-solo-la-nave-pedida.md). La 5 queda sin estado: así sigue el login fallido de las navieras que no son
        # MSC (encargo 45; hasta ahí esta prueba usaba MSC).
        self.assertEqual((n.reservas, {k: (v["estado"], v["detalle"]) for k, v in e["resultados"].items()}),
                         ([], {"7": ("NO ENVIADA", "la fila no trae una nave para reservar")}))
        self.assertTrue(any("No se pudo iniciar sesión en ONE" in l for l in e["lineas"]))

    def test_login_fallido_de_msc_queda_no_enviada(self):
        # Sin la sesión de MSC, cada fila de MSC queda NO ENVIADA con su motivo, sin otro intento de inicio de sesión
        # (decisión de Marcelo, encargo 45, CICLO-login-msc-y-maersk.md); hasta ahí quedaba sin estado.
        with Navieras(self.mod, login_ok=False) as n:
            e = self.correr("MSC", [5, 6])
        self.assertEqual((n.reservas, [l[0] for l in n.logins],
                          {k: (v["estado"], v["detalle"]) for k, v in e["resultados"].items()}),
                         ([], ["msc"], {"5": ("NO ENVIADA", self.mod.MSC_SIN_SESION),
                                        "6": ("NO ENVIADA", "la fila no trae una nave para reservar")}))
        self.assertTrue(any(f"✗ fila 5: NO ENVIADA · {self.mod.MSC_SIN_SESION}" in l for l in e["lineas"]))

    def test_login_de_msc_deja_su_motivo(self):
        # Si el login de MSC dejó su motivo en el Registro (tras el segundo intento: MSC_SIN_SESION_TRAS_DOS,
        # decisión de Marcelo, encargo 50), es el de cada fila de MSC, en el panel y en log.txt (_motivo_sin_sesion).
        motivo = self.mod._msc_motivo_tras_dos(self.mod.MSC_ERROR_TRAS_LA_CLAVE, "")

        def login(page, creds, reg, on_pausa=None):
            reg.sin_sesion = {"msc": motivo}
            return False
        with Navieras(self.mod, login_ok=False) as n:
            self.mod.NAVIERAS["msc"] = ("MSC", login)
            e = self.correr("MSC", [5])
        self.assertEqual((n.reservas, (e["resultados"]["5"]["estado"], e["resultados"]["5"]["detalle"])),
                         ([], ("NO ENVIADA", motivo)))
        self.assertTrue(any(f"✗ fila 5: NO ENVIADA · {motivo}" in l for l in e["lineas"]))

    def test_reservador_que_revienta(self):
        def respuesta(nav, rsv, on_pausa):
            raise RuntimeError("falla falsa del portal")
        with Navieras(self.mod, respuesta=respuesta) as n:
            e = self.correr("COSCO", [5])
        self.assertEqual((e["resultados"]["5"]["estado"], e["resultados"]["5"]["detalle"]),
                         ("ERROR", "falla falsa del portal"))
        self.assertEqual(n.pw.contextos[0].pages[0].capturas, [("error_cosco_rsv_1.png", False)])

    def test_pausa_y_continuar(self):
        def respuesta(nav, rsv, on_pausa):
            on_pausa("resuelve el paso de prueba")
            return ("OK-EJEMPLO", "siguió tras la pausa")
        with Navieras(self.mod, respuesta=respuesta):
            self.srv.json("/api/correr", {"hoja": "HYUNDAI", "usuario": "op_prueba", "filas": [5], "modo": False})
            fin = time.time() + 10
            while time.time() < fin:
                _, e = self.srv.json("/api/estado")
                if e["pausa"]:
                    break
                time.sleep(0.1)
            self.assertEqual((e["pausa"], e["pausa_msg"], e["estado"]),
                             (True, "resuelve el paso de prueba", "Te toca a ti: resuelve el paso en el navegador"))
            self.assertEqual(self.srv.json("/api/continuar", {}), (200, {"ok": True}))
            e = self.srv.esperar_fin()
        self.assertEqual(e["resultados"]["5"]["detalle"], "siguió tras la pausa")
        self.assertFalse(e["pausa"])

    def test_detener(self):
        def respuesta(nav, rsv, on_pausa):
            on_pausa("esperando para detener")
            return ("OK-EJEMPLO", "fila tras detener")
        with Navieras(self.mod, respuesta=respuesta) as n:
            self.srv.json("/api/correr", {"hoja": "MAERSK", "usuario": "op_prueba", "filas": [5, 6],
                                          "modo": False})
            fin = time.time() + 10
            while time.time() < fin and not self.srv.json("/api/estado")[1]["pausa"]:
                time.sleep(0.1)
            self.assertEqual(self.srv.json("/api/detener", {}), (200, {"ok": True}))
            e = self.srv.esperar_fin()
            time.sleep(0.3)
        self.assertEqual(e["estado"], "Detenido.")
        self.assertEqual([r[1] for r in n.reservas], [5])            # la fila 6 ya no se procesa
        self.assertTrue(n.pw.contextos[0].cerrado)
        self.assertTrue(any("Detenido por el operador" in l for l in e["lineas"]))

    def test_modo_login(self):
        self.assertEqual(self.srv.json("/api/login", {"usuario": "op_prueba", "navieras": ["xx"]}),
                         (400, {"error": "elige al menos una naviera"}))
        with Navieras(self.mod) as n:
            t0 = time.time()
            st, r = self.srv.json("/api/login", {"usuario": "op_prueba", "navieras": ["one", "xx", "msc"]})
            e = self.srv.esperar_fin(maximo=20)
            demora = time.time() - t0
        self.assertEqual((st, r), (200, {"ok": True}))
        self.assertEqual([l[0] for l in n.logins], ["one", "msc"])
        self.assertEqual(e["estado"], "Sesiones abiertas.")
        self.assertGreaterEqual(demora, 6)          # deja el navegador abierto 6 s antes de cerrarlo
        self.assertEqual(Path(n.pw.lanzamientos[0]["user_data_dir"]), self.sb / "perfiles" / "op_prueba")


def puertos_seguidos(n=4, intentos=200):
    """El primero de 'n' puertos seguidos libres en 127.0.0.1, entre el 20000 y el 39999. El lanzador prueba los puertos
    que siguen a su base (PUERTOS_DEL_PANEL, encargo 52), y los que reparte el sistema (desde el 49152 en este Windows)
    vienen seguidos: dos pruebas en paralelo, como las del censo, quedarían en puertos vecinos, y el lanzador de una
    podría dar con el panel de la otra (medido: 29 de 29 seguidos)."""
    for _ in range(intentos):
        p = random.randrange(20000, 40000 - n)
        try:
            for q in range(p, p + n):
                with socket.socket() as s:
                    s.bind(("127.0.0.1", q))
        except OSError:
            continue
        return p
    raise AssertionError(f"no encontré {n} puertos seguidos libres")


def puerto_libre(p):
    """Si nadie escucha en 127.0.0.1:'p' (lo enlaza un instante, sin escuchar). De las pruebas, no del programa: corre
    también contra el código de antes del encargo 53."""
    try:
        with socket.socket() as s:
            s.bind(("127.0.0.1", p))
        return True
    except OSError:
        return False


def servidor_falso(puerto, responder):
    """Un servidor HTTP corriente (con SO_REUSEADDR, como cualquier programa escrito con http.server) en
    127.0.0.1:'puerto', que contesta lo que diga responder(método, ruta): (código, cuerpo). Anota cada pedido. Devuelve
    (servidor, pedidos); se apaga con apagar_falso."""
    pedidos = []

    class Falso(http.server.BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def atender(self, metodo):
            pedidos.append((metodo, self.path))
            codigo, cuerpo = responder(metodo, self.path)
            self.send_response(codigo)
            self.send_header("Content-Length", str(len(cuerpo)))
            self.end_headers()
            self.wfile.write(cuerpo)

        def do_GET(self):
            self.atender("GET")

        def do_POST(self):
            self.atender("POST")

    class Servidor(http.server.ThreadingHTTPServer):
        daemon_threads = True

        def handle_error(self, request, client_address):
            pass                                # una pregunta que vence deja su respuesta sin a quién

    srv = Servidor(("127.0.0.1", puerto), Falso)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv, pedidos


def apagar_falso(srv):
    with contextlib.suppress(Exception):
        srv.shutdown()
    srv.server_close()


def responder_como_panel(al_apagar, corriendo=False, emision=False):
    """Las respuestas de un panel ocioso u ocupado, en un modo, sin ser AQUASHIELD: su /api/estado, su /api/config, y
    su /api/apagar, que llama a al_apagar()."""
    def responder(metodo, ruta):
        if metodo == "POST" and ruta == "/api/apagar":
            al_apagar()
            return 200, b'{"ok": true}'
        if ruta.startswith("/api/estado"):
            return 200, json.dumps({"corriendo": corriendo}).encode()
        if ruta == "/api/config":
            return 200, json.dumps({"modo_emision": emision}).encode()
        return 404, b""
    return responder


class TestArranqueDelPanel(unittest.TestCase):
    """Cómo arranca lanzar_web cuando su puerto está tomado. Desde el encargo 52 prueba también los puertos que siguen
    al base: cada prueba usa los de puertos_seguidos. Desde el 53, mira los cuatro antes de tomar uno libre."""

    AVISO_AJENO = "El puerto {} lo tiene otro programa, que no contesta como AQUASHIELD: no lo cerré."

    def setUp(self):
        soporte.DISPAROS.clear()
        self.orig = (subprocess.check_output, subprocess.call)
        self.llamadas = []

    def tearDown(self):
        subprocess.check_output, subprocess.call = self.orig
        disparos = list(soporte.DISPAROS)
        soporte.DISPAROS.clear()
        soporte.apagar_llaves()
        self.assertEqual(disparos, [])

    def _falsos(self, salida_netstat):
        # netstat y taskkill, si el lanzador los llamara, quedan anotados y no corren (hasta el encargo 52 cerraba así
        # a quien tuviera su puerto).
        def check_output(cmd, shell=False):
            self.llamadas.append(("netstat", cmd))
            return salida_netstat.encode()

        def call(cmd, shell=False):
            self.llamadas.append(("call", cmd))
            return 0
        subprocess.check_output, subprocess.call = check_output, call

    def levantar(self, mod, puerto, abrir=False, segundos=25):
        """lanzar_web de 'mod' en un hilo, hasta que levanta su panel o termina. Devuelve (el hilo, el puerto del panel
        o None)."""
        hilo = threading.Thread(target=mod.lanzar_web, kwargs={"puerto": puerto, "abrir": abrir}, daemon=True)
        hilo.start()
        fin = time.time() + segundos
        while time.time() < fin and hilo.is_alive() and mod._SRV_ACTUAL is None:
            time.sleep(0.05)
        return hilo, (mod._SRV_ACTUAL.server_address[1] if mod._SRV_ACTUAL else None)

    def cerrar(self, mod, hilo):
        """Apaga el panel que haya levantado lanzar_web de 'mod', en el puerto que sea, y espera su hilo."""
        if mod._SRV_ACTUAL:
            with contextlib.suppress(Exception):
                mod._SRV_ACTUAL.shutdown()
        hilo.join(timeout=6)

    def test_no_cierra_al_que_ocupa_el_puerto(self):
        # Decisión de Marcelo, encargo 52 (CICLO-puerto-libre-y-modo-al-armar.md): el lanzador no cierra a quien tiene
        # su puerto y no contesta como AQUASHIELD; prueba el siguiente, y lo dice en el registro del panel. Hasta ahí,
        # le pedía a taskkill que cerrara al pid que netstat daba en ese puerto: el 2026-10-05 lo habría hecho con otro
        # programa que escuchaba en el 8765 (medido sin red con un servidor ajeno). Este ocupante acepta la conexión y
        # calla: el lanzador le pregunta dos veces (1 s y PANEL_ESPERA_LARGA s).
        mod, _ = soporte.cargar()
        p = puertos_seguidos()
        ocupante = socket.socket()
        ocupante.bind(("127.0.0.1", p))
        ocupante.listen(5)
        self._falsos(f"  TCP    127.0.0.1:{p}     0.0.0.0:0      LISTENING       424242\n")
        try:
            hilo, puerto = self.levantar(mod, p)
            try:
                self.assertEqual((puerto, self.llamadas), (p + 1, []))
                self.assertIn(self.AVISO_AJENO.format(p), mod._WEB["log"])
            finally:
                self.cerrar(mod, hilo)
        finally:
            ocupante.close()

    def test_otro_programa_que_contesta_no_se_cierra(self):
        # El caso del 2026-10-05: en el 8765 escucha un programa en Python que no es de AQUASHIELD. Aquí, un http.server
        # que contesta 404: el lanzador le pregunta una sola vez (contestó, y no como un panel), no le pide nada más ni
        # lo cierra, y abre el panel en el puerto siguiente; el otro programa sigue contestando.
        mod, _ = soporte.cargar()
        p = puertos_seguidos()
        ajeno, pedidos = servidor_falso(p, lambda metodo, ruta: (404, b""))
        self._falsos(f"  TCP    127.0.0.1:{p}     0.0.0.0:0      LISTENING       424242\n")
        try:
            hilo, puerto = self.levantar(mod, p)
            try:
                self.assertEqual((puerto, self.llamadas, pedidos), (p + 1, [], [("GET", "/api/estado")]))
                self.assertIn(self.AVISO_AJENO.format(p), mod._WEB["log"])
                self.assertEqual(soporte.Servidor(None, p).get("/api/estado"), (404, b""))
            finally:
                self.cerrar(mod, hilo)
        finally:
            apagar_falso(ajeno)

    def test_no_usa_el_puerto_de_quien_contesta_aunque_se_pueda_enlazar(self):
        # Un programa que escucha en todas las direcciones (0.0.0.0, sin SO_EXCLUSIVEADDRUSE) deja enlazar 127.0.0.1 en
        # su puerto, y lo que llega por 127.0.0.1 pasa al que se enlazó después (medido en este Windows, encargo 52): el
        # panel le quitaría a ese programa lo que le llega por ahí. Si contesta, el lanzador no usa su puerto aunque
        # pudiera enlazarlo. Aquí nadie escucha en 0.0.0.0 (le hablaría al firewall): _panel_en dice que en el puerto
        # base contestó otro programa, y el puerto está libre.
        mod, _ = soporte.cargar()
        p = puertos_seguidos()
        real = mod._panel_en
        mod._panel_en = lambda q, espera=1.0: False if q == p else real(q, espera)
        self._falsos("")
        hilo, puerto = self.levantar(mod, p)
        try:
            self.assertEqual((puerto, self.llamadas), (p + 1, []))
        finally:
            self.cerrar(mod, hilo)

    def test_puerto_que_se_suelta_mientras_pregunta(self):
        # Si el puerto se suelta mientras el lanzador le vuelve a preguntar a quien lo tenía (un panel que se estaba
        # apagando, por ejemplo), lo vuelve a probar y lo toma, como hasta hoy (hasta el encargo 52, después de
        # taskkill). Aquí el ocupante acepta sin contestar y suelta el puerto a los 2 s.
        mod, _ = soporte.cargar()
        p = puertos_seguidos()
        ocupante = socket.socket()
        ocupante.bind(("127.0.0.1", p))
        ocupante.listen(5)
        threading.Timer(2.0, ocupante.close).start()
        self._falsos("")
        try:
            hilo, puerto = self.levantar(mod, p)
            try:
                self.assertEqual((puerto, self.llamadas), (p, []))
                self.assertNotIn(self.AVISO_AJENO.format(p), mod._WEB["log"])
            finally:
                self.cerrar(mod, hilo)
        finally:
            ocupante.close()

    def test_panel_en_dice_quien_contesta(self):
        # _panel_en distingue a nadie (None) de otro programa (False; encargo 52) y de un panel de AQUASHIELD ((su
        # estado, su modo)). El lanzador enlaza el puerto solo si nadie contestó.
        mod, _ = soporte.cargar()
        self.assertIsNone(mod._panel_en(puertos_seguidos(1)))                     # nadie escucha
        casos = [((404, b""), False), ((200, b"no es json"), False), ((200, b'{"otra": 1}'), False),
                 ((200, b'[1, 2]'), False), ((200, b'{"corriendo": false}'), ({"corriendo": False}, True))]
        for respuesta, esperado in casos:
            with self.subTest(respuesta=respuesta):
                p = puertos_seguidos(1)
                falso, _ = servidor_falso(p, lambda metodo, ruta, r=respuesta: (
                    r if ruta.startswith("/api/estado") else (200, b'{"modo_emision": true}')))
                try:
                    self.assertEqual(mod._panel_en(p), esperado)
                finally:
                    apagar_falso(falso)
        # Quien acepta, lee el pedido y cierra sin contestar tampoco es un panel; quien acepta y calla es como nadie.
        for cierra, esperado in ((True, False), (False, None)):
            with self.subTest(cierra=cierra):
                p = puertos_seguidos(1)
                oyente = socket.socket()
                oyente.bind(("127.0.0.1", p))
                oyente.listen(5)

                def atender(oyente=oyente, cierra=cierra):
                    with contextlib.suppress(OSError):
                        conexion, _ = oyente.accept()
                        conexion.recv(65536)
                        if cierra:
                            conexion.close()
                        else:
                            time.sleep(2)
                            conexion.close()
                threading.Thread(target=atender, daemon=True).start()
                try:
                    self.assertIs(mod._panel_en(p), esperado)
                finally:
                    oyente.close()

    def _hijo_servidor(self, puerto, corriendo=False, emision=False):
        # Con 'emision', el panel previo está en modo emisión sin abrir ninguna llave: su es_modo_emision dice True
        # (encargo 51).
        sb = soporte.nueva_sandbox()
        codigo = (f"import sys; sys.path.insert(0, {str(TESTS)!r}); import soporte; "
                  f"mod, sb = soporte.cargar(); mod._WEB['corriendo'] = {corriendo}; "
                  f"mod.es_modo_emision = lambda: {emision}; "
                  f"mod.lanzar_web(puerto={puerto}, abrir=False)")
        return soporte._POPEN([sys.executable, "-c", codigo], cwd=str(sb), stdout=subprocess.DEVNULL,
                              stderr=subprocess.DEVNULL, env={**__import__("os").environ, "PYTHONIOENCODING": "utf-8"})

    def _cerrar_hijo(self, hijo, puerto):
        """Apaga el panel hijo por su /api/apagar, para que su proceso termine solo y borre su sandbox; si no termina
        en 10 s, lo mata. Un hijo matado deja su sandbox en el temporal (encargo 41); desde el 51, se apaga antes."""
        if hijo.poll() is None:
            with contextlib.suppress(Exception):
                soporte.Servidor(None, puerto).post("/api/apagar")
            with contextlib.suppress(subprocess.TimeoutExpired):
                hijo.wait(timeout=10)
        if hijo.poll() is None:
            hijo.kill()

    def _esperar(self, srv, segundos=20):
        fin = time.time() + segundos
        while time.time() < fin:
            try:
                return srv.get("/api/estado")
            except Exception:
                time.sleep(0.2)
        raise AssertionError("no respondió")

    def test_instancia_previa_ociosa_se_apaga(self):
        # Le pide que se apague y toma su puerto. Hasta el encargo 52 esperaba 0,6 s y, si el puerto seguía tomado,
        # cerraba con taskkill a quien lo tuviera; ahora espera que lo suelte, hasta PANEL_ESPERA_LARGA s (sin carga,
        # a los 0,52 s): el relevo queda siempre en el mismo puerto.
        self._falsos("")
        p = puertos_seguidos()
        hijo = self._hijo_servidor(p)
        try:
            self._esperar(soporte.Servidor(None, p))
            cfg = sinteticos.config()
            cfg["usuarios"] = {"op_relevo": cfg["usuarios"]["op_prueba"]}
            mod, _ = soporte.cargar(config=cfg)
            hilo = threading.Thread(target=mod.lanzar_web, kwargs={"puerto": p, "abrir": False}, daemon=True)
            hilo.start()
            self.assertEqual(hijo.wait(timeout=15), 0)               # la instancia previa se apagó sola
            srv = soporte.Servidor(mod, p)
            self._esperar(srv, 10)
            self.assertEqual(srv.json("/api/config")[1]["usuarios"], ["op_relevo"])
            self.assertEqual(self.llamadas, [])
            srv.hilo = hilo
            srv.apagar()
        finally:
            self._cerrar_hijo(hijo, p)

    def test_instancia_previa_ocupada_no_se_toca(self):
        self._falsos("")
        p = puertos_seguidos()
        hijo = self._hijo_servidor(p, corriendo=True)
        try:
            previa = soporte.Servidor(None, p)
            self._esperar(previa)
            mod, _ = soporte.cargar()
            hilo = threading.Thread(target=mod.lanzar_web, kwargs={"puerto": p, "abrir": False}, daemon=True)
            hilo.start()
            hilo.join(timeout=5)                                     # vuelve enseguida, sin levantar nada
            vivo = hilo.is_alive()
            if vivo and mod._SRV_ACTUAL:
                mod._SRV_ACTUAL.shutdown()
            self.assertFalse(vivo, "lanzar_web levantó un servidor en vez de ceder a la instancia ocupada")
            self.assertIsNone(hijo.poll())
            self.assertIs(previa.json("/api/estado")[1]["corriendo"], True)
            self.assertIsNone(mod._SRV_ACTUAL)
            self.assertEqual(self.llamadas, [])
        finally:
            self._cerrar_hijo(hijo, p)

    def test_ocioso_que_tarda_en_soltar_el_puerto(self):
        # Al panel ocioso del mismo modo le pide que se apague y toma su puerto, como hasta hoy, aunque tarde en
        # soltarlo: lo espera hasta PANEL_ESPERA_LARGA s. Hasta el encargo 52, a los 0,6 s cerraba con taskkill a todos
        # los que siguieran en el puerto, por los pid que daba netstat, fueran o no de AQUASHIELD (un programa en
        # 0.0.0.0 puede escuchar en el mismo puerto, con el panel encima en 127.0.0.1). Sin carga, un panel lo suelta a
        # los 0,52 s (medido 5 veces); este, a los 1,5 s.
        self._falsos("")
        p = puertos_seguidos()
        caja = {}

        def al_apagar():
            def despues():
                time.sleep(1.5)
                apagar_falso(caja["srv"])
            threading.Thread(target=despues, daemon=True).start()
        caja["srv"], pedidos = servidor_falso(p, responder_como_panel(al_apagar))
        mod, _ = soporte.cargar()
        mod.es_modo_emision = lambda: False
        try:
            hilo, puerto = self.levantar(mod, p)
            try:
                self.assertEqual((puerto, self.llamadas), (p, []))
                self.assertEqual(pedidos, [("GET", "/api/estado"), ("GET", "/api/config"), ("POST", "/api/apagar")])
            finally:
                self.cerrar(mod, hilo)
        finally:
            apagar_falso(caja["srv"])

    def test_ocioso_que_no_suelta_el_puerto_no_se_cierra(self):
        # Si el panel ocioso no suelta su puerto en PANEL_ESPERA_LARGA s, el lanzador no lo cierra: lo dice, y prueba
        # el puerto siguiente (encargo 52). Aquí el plazo es de 1,5 s, para no esperar 5.
        self._falsos("")
        p = puertos_seguidos()
        falso, _ = servidor_falso(p, responder_como_panel(lambda: None))
        mod, _ = soporte.cargar()
        mod.es_modo_emision = lambda: False
        mod.PANEL_ESPERA_LARGA = 1.5
        try:
            hilo, puerto = self.levantar(mod, p)
            try:
                self.assertEqual((puerto, self.llamadas), (p + 1, []))
                self.assertIn(f"El panel de AQUASHIELD del puerto {p} no lo soltó en 1.5 s, después de pedirle que se "
                              f"apagara: no lo cerré.", mod._WEB["log"])
                self.assertEqual(soporte.Servidor(None, p).json("/api/estado"), (200, {"corriendo": False}))
            finally:
                self.cerrar(mod, hilo)
        finally:
            apagar_falso(falso)

    MODO = {True: "EMISIÓN (reservas reales)", False: "prueba (sin emitir)"}

    def texto_otro_modo(self, puerto, modo, propio, ocupado):
        """El aviso del lanzador ante un panel en el otro modo (o que no dice el suyo, con 'modo' None)."""
        cual = f"en modo {self.MODO[modo]}" if modo is not None else "y no pude saber en qué modo está"
        cerrar = "con su botón rojo de apagar, arriba a la derecha, o cerrando su pestaña y esperando unos dos minutos"
        como = (f"Ese panel tiene una corrida en curso: espera a que termine, después ciérralo {cerrar}, y vuelve a "
                f"abrir este lanzador." if ocupado else f"Ciérralo {cerrar}, y vuelve a abrir este lanzador.")
        return (f"Ya hay un panel de AQUASHIELD abierto {cual}, en http://127.0.0.1:{puerto}/, y este lanzador lo "
                f"abriría en modo {self.MODO[propio]}. Para no mezclar los modos, no me conecto a ese panel, no lo "
                f"cierro y no abro otro. {como}")

    def lanzar(self, puerto, propio, abrir=True):
        """Corre lanzar_web de una copia nueva, en modo 'propio' (su es_modo_emision, sin abrir ninguna llave), con su
        aviso anotado. Devuelve (si quedó sirviendo, los avisos, la copia)."""
        mod, _ = soporte.cargar()
        mod.es_modo_emision = (lambda: propio)
        avisos = []
        mod._avisar_al_lanzar = avisos.append
        hilo = threading.Thread(target=mod.lanzar_web, kwargs={"puerto": puerto, "abrir": abrir}, daemon=True)
        hilo.start()
        hilo.join(timeout=10)
        vivo = hilo.is_alive()
        if vivo and mod._SRV_ACTUAL:
            mod._SRV_ACTUAL.shutdown()
            hilo.join(timeout=6)
        return vivo, avisos, mod

    def test_panel_en_otro_modo_avisa_y_no_corre(self):
        # Decisión de Marcelo, encargo 51 (CICLO-modo-y-lanzadores.md): un lanzador nunca se conecta a un panel abierto
        # en otro modo; avisa y no corre. Medido antes del cambio, en estos mismos cuatro casos: con el panel ocupado,
        # el lanzador lo abría en el navegador en el modo que tuviera (también el de prueba a uno de emisión); con el
        # panel ocioso, le pedía que se apagara y tomaba su puerto, y una pestaña suya que quedara abierta hablaba con
        # el panel nuevo, del otro modo. Ahora: un solo aviso; sin abrir el navegador (abrir=True: la trampa lo
        # anotaría), sin apagar ni tocar ese panel, sin netstat ni taskkill, y sin levantar otro.
        self._falsos("")
        for emision in (False, True):
            for ocupado in (False, True):
                with self.subTest(panel="emisión" if emision else "prueba", ocupado=ocupado):
                    p = puertos_seguidos()
                    hijo = self._hijo_servidor(p, corriendo=ocupado, emision=emision)
                    try:
                        previa = soporte.Servidor(None, p)
                        self._esperar(previa)
                        vivo, avisos, mod = self.lanzar(p, not emision)
                        self.assertFalse(vivo, "el lanzador levantó un servidor")
                        self.assertIsNone(mod._SRV_ACTUAL)
                        self.assertEqual(avisos, [self.texto_otro_modo(p, emision, not emision, ocupado)])
                        self.assertIsNone(hijo.poll())
                        self.assertEqual((previa.json("/api/estado")[1]["corriendo"],
                                          previa.json("/api/config")[1]["modo_emision"]), (ocupado, emision))
                        self.assertEqual(self.llamadas, [])
                    finally:
                        self._cerrar_hijo(hijo, p)

    def test_panel_en_el_puerto_siguiente_como_hoy(self):
        # Los paneles de AQUASHIELD se tratan como hasta hoy en el puerto base, también en el que sigue (decisión de
        # Marcelo, encargo 52): con otro programa en el base, un panel en el siguiente, en el otro modo, avisa y no
        # corre; en el mismo y ocupado, no levanta otro (aquí no abre el navegador: abrir=False); en el mismo y ocioso,
        # le pide que se apague y toma su puerto. El lanzador es de prueba.
        self._falsos("")
        for emision, ocupado in ((True, False), (True, True), (False, True), (False, False)):
            with self.subTest(panel="emisión" if emision else "prueba", ocupado=ocupado):
                p = puertos_seguidos()
                ajeno, _ = servidor_falso(p, lambda metodo, ruta: (404, b""))
                hijo = self._hijo_servidor(p + 1, corriendo=ocupado, emision=emision)
                try:
                    self._esperar(soporte.Servidor(None, p + 1))
                    mod, _ = soporte.cargar()
                    mod.es_modo_emision = lambda: False
                    avisos = []
                    mod._avisar_al_lanzar = avisos.append
                    hilo, puerto = self.levantar(mod, p)
                    try:
                        if emision:
                            self.assertEqual((puerto, avisos),
                                             (None, [self.texto_otro_modo(p + 1, True, False, ocupado)]))
                            self.assertIsNone(hijo.poll())
                        elif ocupado:
                            self.assertEqual((puerto, avisos), (None, []))
                            self.assertIsNone(hijo.poll())
                        else:
                            self.assertEqual((puerto, avisos), (p + 1, []))
                            self.assertEqual(hijo.wait(timeout=15), 0)
                        self.assertEqual(self.llamadas, [])
                    finally:
                        self.cerrar(mod, hilo)
                finally:
                    self._cerrar_hijo(hijo, p + 1)
                    apagar_falso(ajeno)

    def test_panel_que_no_dice_su_modo_avisa_y_no_corre(self):
        # Si el panel abierto responde como AQUASHIELD pero no dice su modo (su /api/config falla, o no trae
        # «modo_emision» como true o false), el lanzador tampoco se conecta: avisa que no supo en qué modo está, y no
        # le pide que se apague.
        self._falsos("")
        for config in (None, {"version": "otra"}, {"modo_emision": "false"}):
            with self.subTest(config=config):
                def responder(metodo, ruta, cfg=config):
                    if metodo == "GET" and ruta.startswith("/api/estado"):
                        return 200, json.dumps({"corriendo": False}).encode()
                    if metodo == "GET" and ruta == "/api/config" and cfg is not None:
                        return 200, json.dumps(cfg).encode()
                    return (500, b"falla de prueba") if metodo == "GET" else (200, b"{}")
                p = puertos_seguidos()
                srv, pedidos = servidor_falso(p, responder)
                try:
                    vivo, avisos, mod = self.lanzar(p, False)
                    self.assertEqual((vivo, mod._SRV_ACTUAL), (False, None))
                    self.assertEqual(avisos, [self.texto_otro_modo(p, None, False, False)])
                    self.assertEqual([m for m, _ in pedidos], ["GET", "GET"])
                    self.assertEqual(self.llamadas, [])
                finally:
                    apagar_falso(srv)

    def test_mismo_modo_de_emision_cede_como_antes(self):
        # En el mismo modo, como antes: al panel de emisión ocupado, el lanzador de emisión lo deja como está y no
        # levanta otro, sin aviso (aquí no abre el navegador: abrir=False).
        self._falsos("")
        p = puertos_seguidos()
        hijo = self._hijo_servidor(p, corriendo=True, emision=True)
        try:
            self._esperar(soporte.Servidor(None, p))
            vivo, avisos, mod = self.lanzar(p, True, abrir=False)
            self.assertEqual((vivo, avisos, mod._SRV_ACTUAL, hijo.poll(), self.llamadas), (False, [], None, None, []))
        finally:
            self._cerrar_hijo(hijo, p)

    def test_panel_lento_del_otro_modo_no_se_cierra(self):
        # Un panel del otro modo que tarda más de 1 s en contestar (con la máquina cargada) avisa y no corre: si el
        # puerto sigue tomado, el lanzador le vuelve a preguntar, hasta PANEL_ESPERA_LARGA s (revisión del encargo 51:
        # a 1 s lo daba por ajeno y le pedía a taskkill que lo cerrara; desde el encargo 52, al ajeno no lo cierra, pero
        # abriría el suyo en el puerto siguiente, junto a un panel del otro modo).
        def responder(metodo, ruta):
            if ruta.startswith("/api/estado"):
                time.sleep(1.6)
                return 200, json.dumps({"corriendo": True}).encode()
            return 200, json.dumps({"modo_emision": True}).encode()
        p = puertos_seguidos()
        srv, pedidos = servidor_falso(p, responder)
        try:
            self._falsos(f"  TCP    127.0.0.1:{p}     0.0.0.0:0      LISTENING       424242\n")
            vivo, avisos, mod = self.lanzar(p, False)
            self.assertEqual(mod.PANEL_ESPERA_LARGA, 5.0)
            self.assertEqual((vivo, mod._SRV_ACTUAL, self.llamadas), (False, None, []))
            self.assertEqual(avisos, [self.texto_otro_modo(p, True, False, True)])
            self.assertEqual([r for _, r in pedidos], ["/api/estado", "/api/estado", "/api/config"])
        finally:
            apagar_falso(srv)

    def test_mira_los_cuatro_puertos_con_el_base_libre(self):
        # Decisión de Marcelo, encargo 53 (CICLO-bloqueo-cma-y-sandbox.md): el lanzador mira los cuatro puertos en cada
        # arranque, y un panel de AQUASHIELD en cualquiera de ellos se trata como hasta hoy, también con el base libre.
        # Aquí el panel está en el último de los cuatro, y el base y los otros dos están libres. Hasta el encargo 53,
        # el lanzador tomaba el base sin mirar los demás y abría un panel nuevo junto al que ya estaba (medido sin red
        # con el código de 9f45608, con esta misma prueba). Ahora: en el otro modo, avisa y no corre; en el mismo y
        # ocupado, no levanta otro (aquí no abre el navegador: abrir=False); en el mismo y ocioso, le pide que se apague
        # y toma su puerto, no el base. El lanzador es de prueba.
        self._falsos("")
        mod, _ = soporte.cargar()
        mod.es_modo_emision = lambda: False
        avisos = []
        mod._avisar_al_lanzar = avisos.append
        for emision, ocupado in ((True, False), (True, True), (False, True), (False, False)):
            with self.subTest(panel="emisión" if emision else "prueba", ocupado=ocupado):
                p = puertos_seguidos()
                caja = {}

                def al_apagar(caja=caja):
                    threading.Timer(0.3, apagar_falso, args=(caja["srv"],)).start()
                caja["srv"], pedidos = servidor_falso(p + 3, responder_como_panel(al_apagar, ocupado, emision))
                avisos.clear()
                mod._SRV_ACTUAL = None
                try:
                    hilo, puerto = self.levantar(mod, p)
                    try:
                        if emision:
                            self.assertEqual((puerto, avisos),
                                             (None, [self.texto_otro_modo(p + 3, True, False, ocupado)]))
                        elif ocupado:
                            self.assertEqual((puerto, avisos), (None, []))
                        else:
                            self.assertEqual((puerto, avisos), (p + 3, []))
                        apagar = [] if emision or ocupado else [("POST", "/api/apagar")]
                        self.assertEqual(pedidos, [("GET", "/api/estado"), ("GET", "/api/config")] + apagar)
                        self.assertTrue(puerto_libre(p))                  # el base queda sin tomar
                        self.assertEqual(self.llamadas, [])
                    finally:
                        self.cerrar(mod, hilo)
                finally:
                    apagar_falso(caja["srv"])

    def test_panel_lento_en_otro_puerto_con_el_base_libre(self):
        # Un panel del otro modo que tarda más de 1 s en contestar (con la máquina cargada) se reconoce también fuera
        # del base: si nadie contestó y el puerto está tomado, el lanzador le vuelve a preguntar, hasta
        # PANEL_ESPERA_LARGA s, antes de tomar uno libre (encargo 53; en el base, desde la revisión del encargo 51).
        def responder(metodo, ruta):
            if ruta.startswith("/api/estado"):
                time.sleep(1.6)
                return 200, json.dumps({"corriendo": True}).encode()
            return 200, json.dumps({"modo_emision": True}).encode()
        p = puertos_seguidos()
        srv, pedidos = servidor_falso(p + 1, responder)
        try:
            self._falsos("")
            vivo, avisos, mod = self.lanzar(p, False)
            self.assertEqual((vivo, mod._SRV_ACTUAL, self.llamadas), (False, None, []))
            self.assertEqual(avisos, [self.texto_otro_modo(p + 1, True, False, True)])
            self.assertEqual([r for _, r in pedidos], ["/api/estado", "/api/estado", "/api/config"])
            self.assertTrue(puerto_libre(p))
        finally:
            apagar_falso(srv)

    def test_pregunta_a_los_cuatro_a_la_vez(self):
        # La primera pregunta va a los cuatro puertos a la vez (_quien_esta, encargo 53): preguntarle a un puerto donde
        # nadie escucha tarda 1 s en este equipo (encargo 51), y de a uno serían 4 s en cada arranque con los cuatro
        # libres. Aquí cada pregunta espera a las otras tres (una barrera de 4): de a una, la barrera se rompe.
        mod, _ = soporte.cargar()
        barrera = threading.Barrier(4, timeout=4)
        preguntas, rotas = [], []

        def panel_en(q, espera=1.0):
            preguntas.append((q, espera))
            try:
                barrera.wait()
            except threading.BrokenBarrierError:
                rotas.append(q)
            return None
        mod._panel_en = panel_en
        self._falsos("")
        p = puertos_seguidos()
        hilo, puerto = self.levantar(mod, p)
        try:
            self.assertEqual((puerto, rotas, self.llamadas), (p, [], []))
            self.assertEqual(sorted(preguntas), [(q, 1.0) for q in range(p, p + 4)])
        finally:
            self.cerrar(mod, hilo)

    def test_otro_programa_despues_del_puerto_usado_no_se_avisa(self):
        # El aviso de un puerto que tiene otro programa dice por qué el panel no quedó ahí: va solo para los que quedan
        # antes del puerto que usa (encargo 53). Aquí el base está libre y el siguiente lo tiene un programa que
        # contesta 404: el panel queda en el base, sin aviso, y a ese programa solo se le pregunta una vez.
        mod, _ = soporte.cargar()
        p = puertos_seguidos()
        ajeno, pedidos = servidor_falso(p + 1, lambda metodo, ruta: (404, b""))
        self._falsos("")
        try:
            hilo, puerto = self.levantar(mod, p)
            try:
                self.assertEqual((puerto, self.llamadas, pedidos), (p, [], [("GET", "/api/estado")]))
                self.assertFalse([linea for linea in mod._WEB["log"] if "lo tiene otro programa" in linea])
            finally:
                self.cerrar(mod, hilo)
        finally:
            apagar_falso(ajeno)

    def test_con_dos_paneles_decide_el_primero(self):
        # Con más de un panel de AQUASHIELD en los cuatro puertos, decide el primero, en el orden de los puertos, como
        # hasta hoy decidía el del base (encargo 53). Aquí el base está libre, en el siguiente hay un panel ocioso del
        # mismo modo y en el que sigue, uno del otro modo: al ocioso le pide que se apague y toma su puerto; al del
        # otro modo solo le pregunta, y no avisa.
        self._falsos("")
        p = puertos_seguidos()
        caja = {}

        def al_apagar():
            threading.Timer(0.3, apagar_falso, args=(caja["srv"],)).start()
        caja["srv"], pedidos_ocioso = servidor_falso(p + 1, responder_como_panel(al_apagar))
        otro, pedidos_otro = servidor_falso(p + 2, responder_como_panel(lambda: None, emision=True))
        mod, _ = soporte.cargar()
        mod.es_modo_emision = lambda: False
        avisos = []
        mod._avisar_al_lanzar = avisos.append
        try:
            hilo, puerto = self.levantar(mod, p)
            try:
                self.assertEqual((puerto, avisos, self.llamadas), (p + 1, [], []))
                self.assertEqual(pedidos_ocioso,
                                 [("GET", "/api/estado"), ("GET", "/api/config"), ("POST", "/api/apagar")])
                self.assertEqual(pedidos_otro, [("GET", "/api/estado"), ("GET", "/api/config")])
            finally:
                self.cerrar(mod, hilo)
        finally:
            apagar_falso(caja["srv"])
            apagar_falso(otro)

    def test_puerto_que_se_toma_mientras_mira_los_demas(self):
        # Si el puerto libre que iba a tomar lo toma otro mientras el lanzador miraba los demás, a quien lo tomó se le
        # pregunta quién es, como en la primera vuelta (encargo 53): aquí, un panel ocioso del mismo modo, que se apaga
        # cuando se le pide, y el lanzador toma su puerto. Para simular el instante en que lo tomó, la primera pregunta
        # y la primera mirada al base dicen que estaba libre.
        self._falsos("")
        p = puertos_seguidos()
        caja = {}

        def al_apagar():
            threading.Timer(0.3, apagar_falso, args=(caja["srv"],)).start()
        caja["srv"], pedidos = servidor_falso(p, responder_como_panel(al_apagar))
        mod, _ = soporte.cargar()
        mod.es_modo_emision = lambda: False
        real_panel, real_libre = mod._panel_en, mod._puerto_libre
        primera = {"panel": True, "libre": True}

        def panel_en(q, espera=1.0):
            if q == p and espera == 1.0 and primera["panel"]:
                primera["panel"] = False
                return None
            return real_panel(q, espera)

        def libre(q):
            if q == p and primera["libre"]:
                primera["libre"] = False
                return True
            return real_libre(q)
        mod._panel_en, mod._puerto_libre = panel_en, libre
        try:
            hilo, puerto = self.levantar(mod, p)
            try:
                self.assertEqual((puerto, self.llamadas), (p, []))
                self.assertEqual(pedidos, [("GET", "/api/estado"), ("GET", "/api/config"), ("POST", "/api/apagar")])
            finally:
                self.cerrar(mod, hilo)
        finally:
            apagar_falso(caja["srv"])

    def test_con_los_cuatro_tomados_avisa_de_todos(self):
        # Con los cuatro puertos tomados por otros programas, el lanzador no abre el panel web (main() abre el de
        # escritorio) y dice, de cada uno, que lo tiene otro programa (encargo 53: los avisos van después de mirar los
        # cuatro, y sin un puerto que usar, de todos).
        mod, _ = soporte.cargar()
        p = puertos_seguidos()
        ajenos = [servidor_falso(q, lambda metodo, ruta: (404, b""))[0] for q in range(p, p + 4)]
        self._falsos("")
        try:
            with self.assertRaises(RuntimeError):
                mod.lanzar_web(puerto=p, abrir=False)
            self.assertEqual([linea for linea in mod._WEB["log"] if "lo tiene otro programa" in linea],
                             [self.AVISO_AJENO.format(q) for q in range(p, p + 4)])
            self.assertEqual((mod._SRV_ACTUAL, self.llamadas), (None, []))
        finally:
            for ajeno in ajenos:
                apagar_falso(ajeno)

    def test_el_aviso_va_a_una_ventana(self):
        # Los .bat lanzan con pythonw, sin consola: el aviso del lanzador va por la consola y a una ventana, como el de
        # main() cuando el panel web no abre (encargo 51). Aquí la ventana es una trampa del arnés, que la anota.
        mod, _ = soporte.cargar()
        salida = io.StringIO()
        with contextlib.redirect_stdout(salida):
            mod._avisar_al_lanzar("aviso de prueba")
        disparos = list(soporte.DISPAROS)
        soporte.DISPAROS.clear()
        self.assertEqual(salida.getvalue(), "aviso de prueba\n")
        self.assertEqual(disparos, [("tkinter.messagebox.showwarning", repr(("AQUASHIELD", "aviso de prueba")))])


if __name__ == "__main__":
    unittest.main()
